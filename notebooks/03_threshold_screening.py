import os
import re
import csv
from collections import defaultdict, Counter


# ============================================================
# CONFIGURATION
# ============================================================

MIN_IDENTITY = 30.0
MIN_COVERAGE = 70.0
MAX_EVALUE = 0.05

DOMINANCE_THRESHOLD = 0.70

# This is NOT the required number of BLAST hits.
# It only means that >=2 qualifying hits are preferred
# before calling a taxonomic classification strong.
MIN_HITS_FOR_STRONG_CLASSIFICATION = 2


# ============================================================
# CATEGORY DEFINITIONS
# ============================================================

CATEGORIES = {
    1: "KEEP_ARTHROPOD",
    2: "POSSIBLE_BACTERIAL",
    3: "POSSIBLE_FUNGAL",
    4: "NON_ARTHROPOD_EUKARYOTE",
    5: "OTHER_PROKARYOTE",
    6: "AMBIGUOUS",
    7: "NO_QUALIFYING_HIT",
    8: "UNKNOWN_TAXONOMY",
}


# ============================================================
# PATH FUNCTIONS
# ============================================================

def expand_path(path):

    return os.path.abspath(
        os.path.expanduser(
            path.strip().strip('"')
        )
    )


def check_file_exists(path, description):

    if not os.path.isfile(path):

        print("\nERROR: File not found:")
        print(f"{description}: {path}")

        return False

    return True


# ============================================================
# READ RAW BLAST TSV
# ============================================================

def read_blast_file(path):

    print("\nReading RAW BLAST TSV...")

    rows = []

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="replace"
    ) as file:

        reader = csv.reader(
            file,
            delimiter="\t"
        )

        for line_number, row in enumerate(
            reader,
            start=1
        ):

            if not row:
                continue

            # Skip header
            if row[0].strip().lower() == "qseqid":
                continue

            if len(row) < 10:

                print(
                    f"WARNING: Skipping line {line_number}: "
                    f"only {len(row)} columns."
                )

                continue

            try:

                record = {
                    "qseqid": row[0].strip(),
                    "sseqid": row[1].strip(),
                    "pident": float(row[2]),
                    "length": int(float(row[3])),
                    "qlen": int(float(row[4])),
                    "slen": int(float(row[5])),
                    "qcovs": float(row[6]),
                    "evalue": float(row[7]),
                    "bitscore": float(row[8]),
                    "stitle": row[9].strip(),
                }

                rows.append(record)

            except ValueError:

                print(
                    f"WARNING: Could not parse line "
                    f"{line_number}."
                )

    print(
        f"Raw BLAST rows read: {len(rows):,}"
    )

    return rows


# ============================================================
# BLAST QUALITY FILTER
# ============================================================

def passes_quality_filter(hit):

    return (
        hit["pident"] >= MIN_IDENTITY
        and
        hit["qcovs"] >= MIN_COVERAGE
        and
        hit["evalue"] <= MAX_EVALUE
    )


# ============================================================
# APPLY QUALITY FILTER
# ============================================================

def apply_quality_filter(rows):

    print("\nApplying BLAST quality thresholds...")

    qualifying = []
    failed = []

    for hit in rows:

        if passes_quality_filter(hit):

            qualifying.append(hit)

        else:

            failed.append(hit)

    print(
        f"  Total BLAST hits:       {len(rows):,}"
    )

    print(
        f"  Passing all thresholds: {len(qualifying):,}"
    )

    print(
        f"  Failed thresholds:      {len(failed):,}"
    )

    print("\nThresholds used:")

    print(
        f"  Identity >= {MIN_IDENTITY}%"
    )

    print(
        f"  Coverage >= {MIN_COVERAGE}%"
    )

    print(
        f"  E-value  <= {MAX_EVALUE}"
    )

    return qualifying, failed


# ============================================================
# EXTRACT NCBI TAXID
# ============================================================

def extract_taxid(stitle):

    if not stitle:
        return None

    match = re.search(
        r"\bOX=(\d+)",
        stitle
    )

    if match:

        return int(
            match.group(1)
        )

    return None


# ============================================================
# EXTRACT ORGANISM
# ============================================================

def extract_species(stitle):

    if not stitle:
        return "Unknown"

    match = re.search(
        r"\bOS=([^=]+?)(?=\s+(?:OX|GN|PE|SV|CC|KW|GO)=|$)",
        stitle
    )

    if match:

        species = match.group(1).strip()

        if species:
            return species

    return "Unknown"


# ============================================================
# PARSE rankedlineage.dmp
# ============================================================

def parse_ranked_lineage_line(line):

    line = line.strip()

    if not line:
        return None

    parts = [
        x.strip()
        for x in line.split("|")
    ]

    if len(parts) < 10:
        return None

    try:

        taxid = int(parts[0])

    except ValueError:

        return None

    return {
        "taxid": taxid,
        "scientific_name": parts[1],
        "species": parts[2],
        "genus": parts[3],
        "family": parts[4],
        "order": parts[5],
        "class": parts[6],
        "phylum": parts[7],
        "kingdom": parts[8],
        "superkingdom": parts[9],
    }


# ============================================================
# LOAD ONLY TAXIDS FROM QUALIFYING HITS
# ============================================================

def load_required_taxonomy(
    rankedlineage_path,
    required_taxids
):

    print(
        "\nLoading taxonomy ONLY for qualifying BLAST hits..."
    )

    required_taxids = set(
        required_taxids
    )

    taxonomy = {}

    if not required_taxids:

        print(
            "No qualifying TaxIDs to load."
        )

        return taxonomy

    print(
        f"Qualifying TaxIDs required: "
        f"{len(required_taxids):,}"
    )

    with open(
        rankedlineage_path,
        "r",
        encoding="utf-8",
        errors="replace"
    ) as file:

        for line in file:

            record = parse_ranked_lineage_line(
                line
            )

            if record is None:
                continue

            taxid = record["taxid"]

            if taxid in required_taxids:

                taxonomy[taxid] = record

                if (
                    len(taxonomy)
                    == len(required_taxids)
                ):

                    break

    print(
        f"TaxIDs successfully found: "
        f"{len(taxonomy):,}/"
        f"{len(required_taxids):,}"
    )

    return taxonomy


# ============================================================
# CLASSIFY TAXONOMY
# ============================================================

def classify_taxonomy(tax_record):

    if tax_record is None:

        return "UNKNOWN_TAXONOMY"

    superkingdom = (
        tax_record["superkingdom"]
        .strip()
        .lower()
    )

    kingdom = (
        tax_record["kingdom"]
        .strip()
        .lower()
    )

    phylum = (
        tax_record["phylum"]
        .strip()
        .lower()
    )

    # --------------------------------------------------------
    # BACTERIA
    # --------------------------------------------------------

    if superkingdom == "bacteria":

        return "BACTERIAL"

    # --------------------------------------------------------
    # ARCHAEA
    # --------------------------------------------------------

    if superkingdom == "archaea":

        return "OTHER_PROKARYOTE"

    # --------------------------------------------------------
    # FUNGI
    # --------------------------------------------------------

    if kingdom == "fungi":

        return "FUNGAL"

    # --------------------------------------------------------
    # ARTHROPODA
    # --------------------------------------------------------

    if phylum == "arthropoda":

        return "ARTHROPOD"

    # --------------------------------------------------------
    # OTHER EUKARYOTES
    # --------------------------------------------------------

    if superkingdom == "eukaryota":

        return "NON_ARTHROPOD_EUKARYOTE"

    # --------------------------------------------------------
    # EVERYTHING ELSE
    # --------------------------------------------------------

    return "OTHER"


# ============================================================
# GROUP RAW HITS BY QUERY
# ============================================================

def group_hits_by_query(rows):

    grouped = defaultdict(list)

    for row in rows:

        grouped[row["qseqid"]].append(row)

    return grouped


# ============================================================
# SORT QUALIFYING HITS
# ============================================================

def get_qualifying_hits(hits):

    qualifying = [
        hit
        for hit in hits
        if passes_quality_filter(hit)
    ]

    qualifying.sort(
        key=lambda x: (
            x["evalue"],
            -x["bitscore"],
            -x["pident"],
            -x["qcovs"],
        )
    )

    return qualifying


# ============================================================
# TAXONOMIC SUPPORT
# ============================================================

def calculate_taxonomic_support(
    qualifying_hits,
    taxonomy
):

    counts = Counter()

    classified_hits = []

    for hit in qualifying_hits:

        taxid = extract_taxid(
            hit["stitle"]
        )

        tax_record = (
            taxonomy.get(taxid)
            if taxid is not None
            else None
        )

        tax_class = classify_taxonomy(
            tax_record
        )

        counts[tax_class] += 1

        classified_hits.append({

            "hit": hit,

            "taxid": taxid,

            "taxonomy": tax_record,

            "tax_class": tax_class,

        })

    return counts, classified_hits


# ============================================================
# FINAL TAXONOMIC DECISION
# ============================================================

def make_decision(
    qualifying_hits,
    counts
):

    # IMPORTANT:
    # This uses the ACTUAL number of qualifying hits.
    # It does NOT assume 5, 7, 10, 20, or 40 hits.

    total = len(
        qualifying_hits
    )

    # ========================================================
    # ZERO QUALIFYING HITS
    # ========================================================

    if total == 0:

        return (
            "NO_QUALIFYING_HIT",
            "No BLAST hit passed all three quality thresholds."
        )

    # ========================================================
    # KNOWN TAXONOMY
    # ========================================================

    unknown = counts.get(
        "UNKNOWN_TAXONOMY",
        0
    )

    known = total - unknown

    # --------------------------------------------------------
    # ALL QUALIFYING HITS HAVE UNKNOWN TAXONOMY
    # --------------------------------------------------------

    if known == 0:

        return (
            "UNKNOWN_TAXONOMY",
            "Qualifying BLAST hits were found, "
            "but taxonomy could not be resolved."
        )

    # ========================================================
    # TOO FEW QUALIFYING HITS
    #
    # This does NOT mean the BLAST search had too few hits.
    # It means too few hits actually passed the filters.
    # ========================================================

    if total < MIN_HITS_FOR_STRONG_CLASSIFICATION:

        # If there is only one qualifying hit, we can still
        # report it, but call the classification ambiguous
        # rather than pretending it is strongly supported.

        return (
            "AMBIGUOUS",
            f"Only {total} qualifying BLAST hit(s) "
            f"were available. "
            f"At least {MIN_HITS_FOR_STRONG_CLASSIFICATION} "
            f"qualifying hits are preferred for strong "
            f"taxonomic classification."
        )

    # ========================================================
    # IMPORTANT:
    #
    # Taxonomic dominance should be calculated from KNOWN
    # taxonomy hits, not from UNKNOWN_TAXONOMY hits.
    #
    # Example:
    #
    # 5 qualifying hits
    # 4 Arthropod
    # 1 Unknown
    #
    # Arthropod support = 4 / 4 = 100%
    #
    # NOT 4 / 5 = 80%.
    # ========================================================

    arthropod_fraction = (
        counts.get("ARTHROPOD", 0)
        / known
    )

    bacterial_fraction = (
        counts.get("BACTERIAL", 0)
        / known
    )

    fungal_fraction = (
        counts.get("FUNGAL", 0)
        / known
    )

    non_arthropod_euk_fraction = (
        counts.get(
            "NON_ARTHROPOD_EUKARYOTE",
            0
        )
        / known
    )

    other_prokaryote_fraction = (
        counts.get(
            "OTHER_PROKARYOTE",
            0
        )
        / known
    )

    # ========================================================
    # ARTHROPOD
    # ========================================================

    if (
        arthropod_fraction
        >= DOMINANCE_THRESHOLD
    ):

        return (
            "KEEP_ARTHROPOD",
            f"Arthropod support "
            f"{arthropod_fraction:.1%} "
            f"of taxonomically resolved qualifying hits."
        )

    # ========================================================
    # BACTERIAL
    # ========================================================

    if (
        bacterial_fraction
        >= DOMINANCE_THRESHOLD
    ):

        return (
            "POSSIBLE_BACTERIAL",
            f"Bacterial support "
            f"{bacterial_fraction:.1%} "
            f"of taxonomically resolved qualifying hits."
        )

    # ========================================================
    # FUNGAL
    # ========================================================

    if (
        fungal_fraction
        >= DOMINANCE_THRESHOLD
    ):

        return (
            "POSSIBLE_FUNGAL",
            f"Fungal support "
            f"{fungal_fraction:.1%} "
            f"of taxonomically resolved qualifying hits."
        )

    # ========================================================
    # NON-ARTHROPOD EUKARYOTE
    # ========================================================

    if (
        non_arthropod_euk_fraction
        >= DOMINANCE_THRESHOLD
    ):

        return (
            "NON_ARTHROPOD_EUKARYOTE",
            f"Non-arthropod eukaryote support "
            f"{non_arthropod_euk_fraction:.1%} "
            f"of taxonomically resolved qualifying hits."
        )

    # ========================================================
    # OTHER PROKARYOTE
    # ========================================================

    if (
        other_prokaryote_fraction
        >= DOMINANCE_THRESHOLD
    ):

        return (
            "OTHER_PROKARYOTE",
            f"Other-prokaryote support "
            f"{other_prokaryote_fraction:.1%} "
            f"of taxonomically resolved qualifying hits."
        )

    # ========================================================
    # MIXED / AMBIGUOUS
    # ========================================================

    return (
        "AMBIGUOUS",
        "No taxonomic group reached the "
        f"{DOMINANCE_THRESHOLD:.0%} dominance threshold "
        "among taxonomically resolved qualifying hits."
    )


# ============================================================
# SCREEN ALL PROTEINS
# ============================================================

def screen_proteins(
    raw_rows,
    taxonomy
):

    print(
        "\nGrouping RAW BLAST hits by protein..."
    )

    grouped = group_hits_by_query(
        raw_rows
    )

    print(
        f"Unique proteins: {len(grouped):,}"
    )

    results = []

    print(
        "\nScreening proteins..."
    )

    for query_id, hits in grouped.items():

        # ====================================================
        # COUNT ACTUAL RAW HITS
        # ====================================================

        total_blast_hits = len(hits)

        # ====================================================
        # QUALITY FILTER
        #
        # This automatically uses however many hits actually
        # exist for this protein.
        #
        # It does NOT assume 5 / 7 / 10 / 20 / 40.
        # ====================================================

        qualifying_hits = get_qualifying_hits(
            hits
        )

        actual_qualifying_count = len(
            qualifying_hits
        )

        # ====================================================
        # ZERO QUALIFYING HITS
        # ====================================================

        if not qualifying_hits:

            results.append({

                "gene_id":
                    query_id,

                "best_hit_accession":
                    "",

                "best_hit_species":
                    "",

                "best_hit_taxid":
                    "",

                "best_hit_lineage":
                    "",

                "best_identity_pct":
                    "",

                "best_query_coverage_pct":
                    "",

                "best_evalue":
                    "",

                "best_bitscore":
                    "",

                "qualifying_hits":
                    0,

                "total_blast_hits":
                    total_blast_hits,

                "arthropod_hits":
                    0,

                "bacterial_hits":
                    0,

                "fungal_hits":
                    0,

                "non_arthropod_eukaryote_hits":
                    0,

                "other_prokaryote_hits":
                    0,

                "unknown_taxonomy_hits":
                    0,

                "known_taxonomy_hits":
                    0,

                "arthropod_fraction":
                    0,

                "bacterial_fraction":
                    0,

                "fungal_fraction":
                    0,

                "non_arthropod_eukaryote_fraction":
                    0,

                "other_prokaryote_fraction":
                    0,

                "classification":
                    "NO_QUALIFYING_HIT",

                "decision":
                    "NO_QUALIFYING_HIT",

                "reason":
                    "No BLAST hit passed all three "
                    "quality thresholds.",

                "best_hit_description":
                    "",
            })

            continue

        # ====================================================
        # TAXONOMIC SUPPORT
        # ====================================================

        counts, classified_hits = (
            calculate_taxonomic_support(
                qualifying_hits,
                taxonomy
            )
        )

        # ====================================================
        # FINAL DECISION
        # ====================================================

        category, reason = make_decision(
            qualifying_hits,
            counts
        )

        # ====================================================
        # BEST QUALIFYING HIT
        # ====================================================

        best = qualifying_hits[0]

        best_species = extract_species(
            best["stitle"]
        )

        best_taxid = extract_taxid(
            best["stitle"]
        )

        best_tax = taxonomy.get(
            best_taxid
        )

        if best_tax:

            best_lineage = (
                best_tax["superkingdom"]
                + "; "
                + best_tax["kingdom"]
                + "; "
                + best_tax["phylum"]
                + "; "
                + best_tax["class"]
                + "; "
                + best_tax["order"]
            )

        else:

            best_lineage = ""

        # ====================================================
        # ACTUAL TAXONOMY COUNTS
        # ====================================================

        total = actual_qualifying_count

        unknown_count = counts.get(
            "UNKNOWN_TAXONOMY",
            0
        )

        known_count = total - unknown_count

        # ====================================================
        # FRACTIONS
        #
        # Fractions are based on KNOWN taxonomy.
        # ====================================================

        if known_count > 0:

            arthropod_fraction = (
                counts.get(
                    "ARTHROPOD",
                    0
                )
                / known_count
            )

            bacterial_fraction = (
                counts.get(
                    "BACTERIAL",
                    0
                )
                / known_count
            )

            fungal_fraction = (
                counts.get(
                    "FUNGAL",
                    0
                )
                / known_count
            )

            non_arthropod_euk_fraction = (
                counts.get(
                    "NON_ARTHROPOD_EUKARYOTE",
                    0
                )
                / known_count
            )

            other_prokaryote_fraction = (
                counts.get(
                    "OTHER_PROKARYOTE",
                    0
                )
                / known_count
            )

        else:

            arthropod_fraction = 0
            bacterial_fraction = 0
            fungal_fraction = 0
            non_arthropod_euk_fraction = 0
            other_prokaryote_fraction = 0

        # ====================================================
        # SAVE RESULT
        # ====================================================

        results.append({

            "gene_id":
                query_id,

            "best_hit_accession":
                best["sseqid"],

            "best_hit_species":
                best_species,

            "best_hit_taxid":
                best_taxid,

            "best_hit_lineage":
                best_lineage,

            "best_identity_pct":
                best["pident"],

            "best_query_coverage_pct":
                best["qcovs"],

            "best_evalue":
                best["evalue"],

            "best_bitscore":
                best["bitscore"],

            # ACTUAL number of qualifying hits
            "qualifying_hits":
                total,

            # ACTUAL number of raw hits
            "total_blast_hits":
                total_blast_hits,

            "arthropod_hits":
                counts.get(
                    "ARTHROPOD",
                    0
                ),

            "bacterial_hits":
                counts.get(
                    "BACTERIAL",
                    0
                ),

            "fungal_hits":
                counts.get(
                    "FUNGAL",
                    0
                ),

            "non_arthropod_eukaryote_hits":
                counts.get(
                    "NON_ARTHROPOD_EUKARYOTE",
                    0
                ),

            "other_prokaryote_hits":
                counts.get(
                    "OTHER_PROKARYOTE",
                    0
                ),

            "unknown_taxonomy_hits":
                unknown_count,

            "known_taxonomy_hits":
                known_count,

            "arthropod_fraction":
                arthropod_fraction,

            "bacterial_fraction":
                bacterial_fraction,

            "fungal_fraction":
                fungal_fraction,

            "non_arthropod_eukaryote_fraction":
                non_arthropod_euk_fraction,

            "other_prokaryote_fraction":
                other_prokaryote_fraction,

            "classification":
                category,

            "decision":
                category,

            "reason":
                reason,

            "best_hit_description":
                best["stitle"],
        })

    return results


# ============================================================
# WRITE TSV
# ============================================================

def save_tsv(
    rows,
    output_path
):

    if not rows:
        return

    fields = list(
        rows[0].keys()
    )

    with open(
        output_path,
        "w",
        newline="",
        encoding="utf-8"
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=fields,
            delimiter="\t"
        )

        writer.writeheader()

        writer.writerows(rows)

    print(
        f"Saved TSV: {output_path}"
    )


# ============================================================
# WRITE EXCEL
# ============================================================

def save_excel(
    results,
    output_path
):

    try:

        from openpyxl import Workbook
        from openpyxl.styles import Font

    except ImportError:

        print(
            "\nWARNING: openpyxl is not installed."
        )

        print(
            "Install with:"
        )

        print(
            "pip install openpyxl"
        )

        return

    if not results:

        print(
            "No results to write to Excel."
        )

        return

    workbook = Workbook()

    # ========================================================
    # ALL RESULTS
    # ========================================================

    sheet = workbook.active

    sheet.title = "All Results"

    fields = list(
        results[0].keys()
    )

    for col, field in enumerate(
        fields,
        start=1
    ):

        cell = sheet.cell(
            row=1,
            column=col,
            value=field
        )

        cell.font = Font(
            bold=True
        )

    for row_number, record in enumerate(
        results,
        start=2
    ):

        for col, field in enumerate(
            fields,
            start=1
        ):

            sheet.cell(
                row=row_number,
                column=col,
                value=record[field]
            )

    sheet.freeze_panes = "A2"

    sheet.auto_filter.ref = (
        sheet.dimensions
    )

    # ========================================================
    # CATEGORY SHEETS
    # ========================================================

    categories = sorted(
        set(
            row["classification"]
            for row in results
        )
    )

    for category in categories:

        ws = workbook.create_sheet(
            title=category[:31]
        )

        category_rows = [
            row
            for row in results
            if row["classification"]
            == category
        ]

        for col, field in enumerate(
            fields,
            start=1
        ):

            cell = ws.cell(
                row=1,
                column=col,
                value=field
            )

            cell.font = Font(
                bold=True
            )

        for row_number, record in enumerate(
            category_rows,
            start=2
        ):

            for col, field in enumerate(
                fields,
                start=1
            ):

                ws.cell(
                    row=row_number,
                    column=col,
                    value=record[field]
                )

        ws.freeze_panes = "A2"

        ws.auto_filter.ref = (
            ws.dimensions
        )

    workbook.save(
        output_path
    )

    print(
        f"Saved Excel: {output_path}"
    )


# ============================================================
# SAVE CATEGORY TSV FILES
# ============================================================

def save_category_files(
    results,
    output_dir
):

    categories = sorted(
        set(
            row["classification"]
            for row in results
        )
    )

    for category in categories:

        category_rows = [
            row
            for row in results
            if row["classification"]
            == category
        ]

        filename = (
            category
            + ".tsv"
        )

        output_path = os.path.join(
            output_dir,
            filename
        )

        save_tsv(
            category_rows,
            output_path
        )


# ============================================================
# FASTA READER
# ============================================================

def read_fasta(path):

    sequences = {}

    current_id = None
    sequence_parts = []

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="replace"
    ) as file:

        for line in file:

            line = line.strip()

            if not line:
                continue

            if line.startswith(">"):

                if current_id is not None:

                    sequences[current_id] = (
                        "".join(sequence_parts)
                    )

                header = line[1:]

                current_id = (
                    header.split()[0]
                )

                sequence_parts = []

            else:

                sequence_parts.append(
                    line
                )

    if current_id is not None:

        sequences[current_id] = (
            "".join(sequence_parts)
        )

    return sequences


# ============================================================
# FASTA WRITER
# ============================================================

def write_fasta(
    sequence_ids,
    sequences,
    output_path
):

    written = 0
    missing = []

    with open(
        output_path,
        "w",
        encoding="utf-8"
    ) as file:

        for sequence_id in sequence_ids:

            if sequence_id not in sequences:

                missing.append(
                    sequence_id
                )

                continue

            file.write(
                ">"
                + sequence_id
                + "\n"
            )

            seq = sequences[
                sequence_id
            ]

            for i in range(
                0,
                len(seq),
                80
            ):

                file.write(
                    seq[i:i + 80]
                    + "\n"
                )

            written += 1

    print(
        f"\nFASTA sequences written: "
        f"{written:,}"
    )

    if missing:

        print(
            f"WARNING: {len(missing):,} "
            "IDs were not found in FASTA."
        )

        missing_file = (
            output_path
            + ".missing.txt"
        )

        with open(
            missing_file,
            "w",
            encoding="utf-8"
        ) as file:

            for sequence_id in missing:

                file.write(
                    sequence_id
                    + "\n"
                )


# ============================================================
# CATEGORY SELECTION
# ============================================================

def ask_categories():

    print(
        "\n"
        + "=" * 70
    )

    print(
        "FASTA CATEGORY SELECTION"
    )

    print(
        "=" * 70
    )

    for number, category in CATEGORIES.items():

        print(
            f"{number} = {category}"
        )

    print(
        "\nSelect one or multiple categories."
    )

    print(
        "Example: 1"
    )

    print(
        "Example: 1,6"
    )

    print(
        "Example: 2,3"
    )

    while True:

        answer = input(
            "\nEnter category numbers: "
        ).strip()

        if not answer:

            print(
                "Please enter at least one category."
            )

            continue

        try:

            numbers = [
                int(x.strip())
                for x in answer.split(",")
                if x.strip()
            ]

        except ValueError:

            print(
                "Invalid input. Use numbers such as 1,6."
            )

            continue

        invalid = [
            n
            for n in numbers
            if n not in CATEGORIES
        ]

        if invalid:

            print(
                "Invalid category number(s): "
                + ", ".join(
                    str(x)
                    for x in invalid
                )
            )

            continue

        numbers = list(
            dict.fromkeys(
                numbers
            )
        )

        selected = [
            CATEGORIES[n]
            for n in numbers
        ]

        return selected


# ============================================================
# EXPORT SELECTED FASTA
# ============================================================

def export_selected_fasta(
    results,
    fasta_sequences,
    categories,
    output_dir,
    output_prefix
):

    selected_rows = [
        row
        for row in results
        if row["classification"]
        in categories
    ]

    selected_ids = [
        row["gene_id"]
        for row in selected_rows
    ]

    if not selected_ids:

        print(
            "\nNo sequences belong to "
            "the selected categories."
        )

        return

    category_text = "_".join(
        categories
    )

    category_text = re.sub(
        r"[^A-Za-z0-9_.-]+",
        "_",
        category_text
    )

    output_path = os.path.join(
        output_dir,
        f"{output_prefix}_{category_text}.fasta"
    )

    write_fasta(
        selected_ids,
        fasta_sequences,
        output_path
    )

    table_path = os.path.join(
        output_dir,
        f"{output_prefix}_{category_text}.tsv"
    )

    save_tsv(
        selected_rows,
        table_path
    )

    print(
        "\nSelected categories:"
    )

    for category in categories:

        count = sum(
            1
            for row in results
            if row["classification"]
            == category
        )

        print(
            f"  {category}: {count:,}"
        )

    print(
        "\nFASTA saved:"
    )

    print(
        output_path
    )


# ============================================================
# SUMMARY
# ============================================================

def print_summary(results):

    counts = Counter(
        row["classification"]
        for row in results
    )

    print(
        "\n"
        + "=" * 70
    )

    print(
        "CONTAMINATION SCREEN SUMMARY"
    )

    print(
        "=" * 70
    )

    print(
        f"Proteins screened: {len(results):,}"
    )

    for category in [
        "KEEP_ARTHROPOD",
        "POSSIBLE_BACTERIAL",
        "POSSIBLE_FUNGAL",
        "NON_ARTHROPOD_EUKARYOTE",
        "OTHER_PROKARYOTE",
        "AMBIGUOUS",
        "NO_QUALIFYING_HIT",
        "UNKNOWN_TAXONOMY",
    ]:

        count = counts.get(
            category,
            0
        )

        percentage = (
            count / len(results) * 100
            if results
            else 0
        )

        print(
            f"{category:30s} "
            f"{count:8,} "
            f"({percentage:6.2f}%)"
        )


# ============================================================
# OUTPUT NAME
# ============================================================

def ask_output_name():

    print(
        "\n"
        + "=" * 70
    )

    print(
        "OUTPUT NAME"
    )

    print(
        "=" * 70
    )

    print(
        "The output folder will be created beside the RAW BLAST file."
    )

    print(
        "Enter a name for this screening run."
    )

    print(
        "Example: metisa_screen_01"
    )

    while True:

        name = input(
            "\nOutput name: "
        ).strip()

        name = re.sub(
            r'[<>:"/\\|?*]',
            "_",
            name
        )

        name = name.strip(
            ". "
        )

        if not name:

            print(
                "Please enter an output name."
            )

            continue

        return name


# ============================================================
# MAIN
# ============================================================

def main():

    print(
        "\n"
        + "=" * 70
    )

    print(
        "BLAST CONTAMINATION SCREEN"
    )

    print(
        "=" * 70
    )

    print(
        "\nWORKFLOW:"
    )

    print(
        "1. Read RAW BLAST TSV"
    )

    print(
        "2. Count the actual BLAST hits for each protein"
    )

    print(
        "3. Apply identity / coverage / E-value thresholds"
    )

    print(
        "4. Use ALL qualifying hits actually present"
    )

    print(
        "5. Taxonomy screening ONLY on qualifying hits"
    )

    print(
        "6. Generate TSV + Excel + selectable FASTA"
    )

    print(
        "\nIMPORTANT:"
    )

    print(
        "The script does NOT assume a fixed number of BLAST hits."
    )

    print(
        "It works whether each protein has 1, 5, 7, 10, 20, 40, "
        "or another number of hits."
    )

    print(
        "\nTaxonomic categories:"
    )

    print(
        "  KEEP_ARTHROPOD"
    )

    print(
        "  POSSIBLE_BACTERIAL"
    )

    print(
        "  POSSIBLE_FUNGAL"
    )

    print(
        "  NON_ARTHROPOD_EUKARYOTE"
    )

    print(
        "  OTHER_PROKARYOTE"
    )

    print(
        "  AMBIGUOUS"
    )

    print(
        "  NO_QUALIFYING_HIT"
    )

    print(
        "  UNKNOWN_TAXONOMY"
    )

    print(
        "\nThresholds:"
    )

    print(
        f"  Identity  >= {MIN_IDENTITY}%"
    )

    print(
        f"  Coverage  >= {MIN_COVERAGE}%"
    )

    print(
        f"  E-value   <= {MAX_EVALUE}"
    )

    print(
        f"  Taxonomic dominance >= "
        f"{DOMINANCE_THRESHOLD:.0%}"
    )

    print(
        f"  Preferred minimum qualifying hits = "
        f"{MIN_HITS_FOR_STRONG_CLASSIFICATION}"
    )

    # ========================================================
    # RAW BLAST INPUT
    # ========================================================

    blast_file = expand_path(
        input(
            "\nPath to RAW BLAST TSV: "
        )
    )

    if not check_file_exists(
        blast_file,
        "RAW BLAST TSV"
    ):

        return

    # ========================================================
    # ORIGINAL FASTA
    # ========================================================

    fasta_file = expand_path(
        input(
            "\nPath to original protein FASTA: "
        )
    )

    if not check_file_exists(
        fasta_file,
        "Protein FASTA"
    ):

        return

    # ========================================================
    # RANKED LINEAGE
    # ========================================================

    rankedlineage_file = expand_path(
        input(
            "\nPath to rankedlineage.dmp: "
        )
    )

    if not check_file_exists(
        rankedlineage_file,
        "rankedlineage.dmp"
    ):

        return

    # ========================================================
    # OUTPUT NAME
    # ========================================================

    output_prefix = ask_output_name()

    blast_dir = os.path.dirname(
        blast_file
    )

    output_dir = os.path.join(
        blast_dir,
        output_prefix
    )

    if os.path.exists(output_dir):

        print(
            "\nWARNING:"
        )

        print(
            "Output folder already exists:"
        )

        print(
            output_dir
        )

        answer = input(
            "\nContinue and overwrite files in this folder? "
            "(y/n): "
        ).strip().lower()

        if answer != "y":

            print(
                "Cancelled."
            )

            return

    os.makedirs(
        output_dir,
        exist_ok=True
    )

    fasta_output_dir = os.path.join(
        output_dir,
        "fasta_exports"
    )

    os.makedirs(
        fasta_output_dir,
        exist_ok=True
    )

    # ========================================================
    # READ RAW BLAST
    # ========================================================

    raw_rows = read_blast_file(
        blast_file
    )

    if not raw_rows:

        print(
            "\nERROR: No BLAST rows found."
        )

        return

    # ========================================================
    # COUNT RAW HITS
    # ========================================================

    raw_grouped = group_hits_by_query(
        raw_rows
    )

    print(
        "\nActual BLAST hit distribution:"
    )

    hit_counts = Counter(
        len(hits)
        for hits in raw_grouped.values()
    )

    for number_of_hits in sorted(
        hit_counts
    ):

        protein_count = hit_counts[
            number_of_hits
        ]

        print(
            f"  {number_of_hits:>4} hit(s): "
            f"{protein_count:,} protein(s)"
        )

    # ========================================================
    # APPLY QUALITY FILTER
    # ========================================================

    qualifying_rows, failed_rows = (
        apply_quality_filter(
            raw_rows
        )
    )

    # ========================================================
    # SAVE FILTERED BLAST
    # ========================================================

    filtered_blast_path = os.path.join(
        output_dir,
        f"{output_prefix}_qualifying_hits.tsv"
    )

    save_tsv(
        qualifying_rows,
        filtered_blast_path
    )

    # ========================================================
    # SAVE FAILED BLAST HITS
    # ========================================================

    failed_blast_path = os.path.join(
        output_dir,
        f"{output_prefix}_failed_threshold_hits.tsv"
    )

    save_tsv(
        failed_rows,
        failed_blast_path
    )

    # ========================================================
    # GET TAXIDS ONLY FROM QUALIFYING HITS
    # ========================================================

    print(
        "\nExtracting TaxIDs ONLY from "
        "quality-passing hits..."
    )

    required_taxids = set()

    for row in qualifying_rows:

        taxid = extract_taxid(
            row["stitle"]
        )

        if taxid is not None:

            required_taxids.add(
                taxid
            )

    print(
        f"Unique qualifying TaxIDs: "
        f"{len(required_taxids):,}"
    )

    # ========================================================
    # TAXONOMY
    #
    # Failed hits NEVER reach this step.
    # ========================================================

    taxonomy = load_required_taxonomy(
        rankedlineage_file,
        required_taxids
    )

    # ========================================================
    # SCREEN ALL PROTEINS
    # ========================================================

    results = screen_proteins(
        raw_rows,
        taxonomy
    )

    # ========================================================
    # SAVE ALL RESULTS TSV
    # ========================================================

    all_tsv = os.path.join(
        output_dir,
        f"{output_prefix}_screening_all.tsv"
    )

    save_tsv(
        results,
        all_tsv
    )

    # ========================================================
    # SAVE CATEGORY TSVs
    # ========================================================

    save_category_files(
        results,
        output_dir
    )

    # ========================================================
    # SAVE EXCEL
    # ========================================================

    excel_file = os.path.join(
        output_dir,
        f"{output_prefix}_screening.xlsx"
    )

    save_excel(
        results,
        excel_file
    )

    # ========================================================
    # SUMMARY
    # ========================================================

    print_summary(
        results
    )

    # ========================================================
    # READ ORIGINAL FASTA
    # ========================================================

    print(
        "\nReading original protein FASTA..."
    )

    fasta_sequences = read_fasta(
        fasta_file
    )

    print(
        f"FASTA sequences loaded: "
        f"{len(fasta_sequences):,}"
    )

    # ========================================================
    # ASK FASTA CATEGORIES
    # ========================================================

    categories = ask_categories()

    # ========================================================
    # EXPORT SELECTED FASTA
    # ========================================================

    export_selected_fasta(
        results,
        fasta_sequences,
        categories,
        fasta_output_dir,
        output_prefix
    )

    # ========================================================
    # FINAL
    # ========================================================

    print(
        "\n"
        + "=" * 70
    )

    print(
        "SCREENING COMPLETE"
    )

    print(
        "=" * 70
    )

    print(
        "\nOutput folder:"
    )

    print(
        output_dir
    )

    print(
        "\nImportant:"
    )

    print(
        "The number of BLAST hits is determined from the actual "
        "input file."
    )

    print(
        "No fixed hit count such as 5, 7, 10, 20, or 40 is assumed."
    )

    print(
        "Only hits passing identity + coverage + E-value "
        "are used for taxonomy."
    )

    print(
        "Unknown-taxonomy hits are excluded from the denominator "
        "when calculating taxonomic dominance."
    )

    print(
        "\nThe selected FASTA can be used for "
        "downstream annotation such as InterProScan."
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    main()
