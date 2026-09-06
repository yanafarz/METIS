import os
import re
import csv
from collections import defaultdict, Counter

MIN_IDENTITY = 30.0
MIN_COVERAGE = 70.0
MAX_EVALUE = 0.05
DOMINANCE_THRESHOLD = 0.70

CATEGORIES = {
    1: "KEEP_ARTHROPOD",
    2: "POSSIBLE_BACTERIAL",
    3: "POSSIBLE_FUNGAL",
    4: "NON_ARTHROPOD_EUKARYOTE",
    5: "OTHER_PROKARYOTE",
    6: "AMBIGUOUS",
    7: "NO_QUALIFYING_HIT",
    8: "UNKNOWN_TAXONOMY",
}

def expand_path(path):
    return os.path.abspath(os.path.expanduser(path.strip().strip('"')))

def check_file_exists(path, description):
    if not os.path.isfile(path):
        print("\nERROR: File not found:")
        print(f"{description}: {path}")
        return False
    return True

def read_blast_file(path):
    print("\nReading RAW BLAST TSV...")
    rows = []

    with open(path, "r", encoding="utf-8", errors="replace") as file:
        reader = csv.reader(file, delimiter="\t")

        for line_number, row in enumerate(reader, start=1):
            if not row:
                continue

            # Skip header
            if row[0].strip().lower() == "qseqid":
                continue

            if len(row) < 10:
                print(f"WARNING: Skipping line {line_number}: only {len(row)} columns.")
                continue

            try:
                record = {
                    "qseqid": row[0].strip(),
                    "sseqid": row[1].strip(),
                    "pident": float(row[2]),
                    "length": int(float(row[3])),
                    "qlen": int(float(row[4])),
                    "slen": int(float(row[5])),
                    "qcov": float(row[6]),
                    "evalue": float(row[7]),
                    "bitscore": float(row[8]),
                    "stitle": row[9].strip(),
                }
                rows.append(record)

            except ValueError:
                print(f"WARNING: Could not parse line {line_number}.")

    print(f"Raw BLAST rows read: {len(rows):,}")
    return rows

def passes_quality_filter(hit):
    return hit["pident"] >= MIN_IDENTITY and hit["qcov"] >= MIN_COVERAGE and hit["evalue"] <= MAX_EVALUE

def apply_quality_filter(rows):
    print("\nApplying BLAST quality thresholds...")
    qualifying, failed = [], []

    for hit in rows:
        (qualifying if passes_quality_filter(hit) else failed).append(hit)

    print(f"  Total BLAST hits:       {len(rows):,}")
    print(f"  Passing all thresholds: {len(qualifying):,}")
    print(f"  Failed thresholds:      {len(failed):,}")

    print("\nThresholds used:")
    print(f"  Identity >= {MIN_IDENTITY}%")
    print(f"  Coverage >= {MIN_COVERAGE}%")
    print(f"  E-value  <= {MAX_EVALUE}")

    return qualifying, failed

def extract_taxid(stitle):
    if not stitle:
        return None

    match = re.search(r"\bOX=(\d+)", stitle)
    if match:
        return int(match.group(1))

    return None

def extract_species(stitle):
    if not stitle:
        return "Unknown"

    match = re.search(r"\bOS=([^=]+?)(?=\s+(?:OX|GN|PE|SV|CC|KW|GO)=|$)", stitle)

    if match:
        species = match.group(1).strip()
        if species:
            return species

    return "Unknown"

def parse_ranked_lineage_line(line):
    line = line.strip()
    if not line:
        return None

    parts = [x.strip() for x in line.split("|")]
    if len(parts) < 10:
        return None

    try:
        taxid = int(parts[0])
    except ValueError:
        return None

    return {
        "taxid": taxid,
        "scientific_name": parts[1],
        "species": parts[2],
        "genus": parts[3],
        "family": parts[4],
        "order": parts[5],
        "class": parts[6],
        "phylum": parts[7],
        "kingdom": parts[8],
        "superkingdom": parts[9],
    }

def load_required_taxonomy(rankedlineage_path, required_taxids):
    print("\nLoading taxonomy ONLY for qualifying BLAST hits...")

    required_taxids = set(required_taxids)
    taxonomy = {}

    if not required_taxids:
        print("No qualifying TaxIDs to load.")
        return taxonomy

    print(f"Qualifying TaxIDs required: {len(required_taxids):,}")

    with open(rankedlineage_path, "r", encoding="utf-8", errors="replace") as file:
        for line in file:
            record = parse_ranked_lineage_line(line)
            if record is None:
                continue

            taxid = record["taxid"]

            if taxid in required_taxids:
                taxonomy[taxid] = record

                if len(taxonomy) == len(required_taxids):
                    break

    print(f"TaxIDs successfully found: {len(taxonomy):,}/{len(required_taxids):,}")

    return taxonomy

def classify_taxonomy(tax_record):
    if tax_record is None:
        return "UNKNOWN_TAXONOMY"

    superkingdom = tax_record["superkingdom"].strip().lower()
    kingdom = tax_record["kingdom"].strip().lower()
    phylum = tax_record["phylum"].strip().lower()

    if superkingdom == "bacteria":
        return "BACTERIAL"
    
    if superkingdom == "archaea":
        return "OTHER_PROKARYOTE"

    if kingdom == "fungi":
        return "FUNGAL"

    if phylum == "arthropoda":
        return "ARTHROPOD"

    if superkingdom == "eukaryota":
        return "NON_ARTHROPOD_EUKARYOTE"

    return "OTHER"

def group_hits_by_query(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["qseqid"]].append(row)
    return grouped

def get_qualifying_hits(hits):
    qualifying = [hit for hit in hits if passes_quality_filter(hit)]

    qualifying.sort(
        key=lambda x: (
            x["evalue"],
            -x["bitscore"],
            -x["pident"],
            -x["qcov"],
        )
    )

    return qualifying

def calculate_taxonomic_support(qualifying_hits, taxonomy):
    counts = Counter()
    classified_hits = []

    for hit in qualifying_hits:
        taxid = extract_taxid(hit["stitle"])
        tax_record = taxonomy.get(taxid) if taxid is not None else None
        tax_class = classify_taxonomy(tax_record)
        counts[tax_class] += 1

        classified_hits.append({
            "hit": hit,
            "taxid": taxid,
            "taxonomy": tax_record,
            "tax_class": tax_class,
        })

    return counts, classified_hits

def make_decision(qualifying_hits, counts):
    total = len(qualifying_hits)

    if total == 0:
        return (
            "NO_QUALIFYING_HIT",
            "No BLAST hit passed all three quality thresholds."
        )

    arthropod = counts.get("ARTHROPOD", 0)
    bacterial = counts.get("BACTERIAL", 0)
    fungal = counts.get("FUNGAL", 0)
    non_arthropod_eukaryote = counts.get("NON_ARTHROPOD_EUKARYOTE", 0)
    other_prokaryote = counts.get("OTHER_PROKARYOTE", 0)
    unknown = counts.get("UNKNOWN_TAXONOMY", 0)

    arthropod_fraction = arthropod / total
    bacterial_fraction = bacterial / total
    fungal_fraction = fungal / total
    non_arthropod_eukaryote_fraction = non_arthropod_eukaryote / total
    other_prokaryote_fraction = other_prokaryote / total
    unknown_fraction = unknown / total

    group_counts = {
        "KEEP_ARTHROPOD": arthropod,
        "POSSIBLE_BACTERIAL": bacterial,
        "POSSIBLE_FUNGAL": fungal,
        "NON_ARTHROPOD_EUKARYOTE": non_arthropod_eukaryote,
        "OTHER_PROKARYOTE": other_prokaryote,
    }

    highest_count = max(group_counts.values())

    if highest_count == 0:
        return (
            "UNKNOWN_TAXONOMY",
            "All qualifying BLAST hits had unresolved taxonomy."
        )

    highest_groups = [
        category
        for category, count in group_counts.items()
        if count == highest_count
    ]

    if len(highest_groups) > 1:
        tied_groups = ", ".join(highest_groups)
        return (
            "AMBIGUOUS",
            f"Taxonomic evidence is tied between {tied_groups} "
            f"({highest_count} qualifying hit(s) each)."
        )

    winner = highest_groups[0]
    winner_fraction = highest_count / total

    if winner_fraction >= DOMINANCE_THRESHOLD:
        if winner == "KEEP_ARTHROPOD":
            return (
                "KEEP_ARTHROPOD",
                f"Arthropod is the dominant taxonomic group "
                f"({highest_count}/{total} = {winner_fraction:.1%} of all qualifying hits)."
            )

        if winner == "POSSIBLE_BACTERIAL":
            return (
                "POSSIBLE_BACTERIAL",
                f"Bacteria is the dominant taxonomic group "
                f"({highest_count}/{total} = {winner_fraction:.1%} of all qualifying hits)."
            )

        if winner == "POSSIBLE_FUNGAL":
            return (
                "POSSIBLE_FUNGAL",
                f"Fungi is the dominant taxonomic group "
                f"({highest_count}/{total} = {winner_fraction:.1%} of all qualifying hits)."
            )

        if winner == "NON_ARTHROPOD_EUKARYOTE":
            return (
                "NON_ARTHROPOD_EUKARYOTE",
                f"Non-arthropod eukaryotes are the dominant taxonomic group "
                f"({highest_count}/{total} = {winner_fraction:.1%} of all qualifying hits)."
            )

        if winner == "OTHER_PROKARYOTE":
            return (
                "OTHER_PROKARYOTE",
                f"Other prokaryotes are the dominant taxonomic group "
                f"({highest_count}/{total} = {winner_fraction:.1%} of all qualifying hits)."
            )

    return (
        "AMBIGUOUS",
        f"{winner} has the highest support "
        f"({highest_count}/{total} = {winner_fraction:.1%}), "
        f"but does not reach the {DOMINANCE_THRESHOLD:.0%} dominance threshold."
    )

def screen_proteins(raw_rows, taxonomy):

    print("\nGrouping RAW BLAST hits by protein...")
    grouped = group_hits_by_query(raw_rows)
    print(f"Unique proteins: {len(grouped):,}")

    results = []

    print("\nScreening proteins...")

    for query_id, hits in grouped.items():

        total_blast_hits = len(hits)

        qualifying_hits = get_qualifying_hits(hits)
        actual_qualifying_count = len(qualifying_hits)

        if not qualifying_hits:

            results.append({
                "gene_id": query_id,
                "best_hit_accession": "",
                "best_hit_species": "",
                "best_hit_taxid": "",
                "best_hit_lineage": "",
                "best_identity_pct": "",
                "best_query_coverage_pct": "",
                "best_evalue": "",
                "best_bitscore": "",
                "qualifying_hits": 0,
                "total_blast_hits": total_blast_hits,
                "arthropod_hits": 0,
                "bacterial_hits": 0,
                "fungal_hits": 0,
                "non_arthropod_eukaryote_hits": 0,
                "other_prokaryote_hits": 0,
                "unknown_taxonomy_hits": 0,
                "known_taxonomy_hits": 0,
                "arthropod_fraction": 0,
                "bacterial_fraction": 0,
                "fungal_fraction": 0,
                "non_arthropod_eukaryote_fraction": 0,
                "other_prokaryote_fraction": 0,
                "classification": "NO_QUALIFYING_HIT",
                "decision": "NO_QUALIFYING_HIT",
                "reason": "No BLAST hit passed all three quality thresholds.",
                "best_hit_description": "",
            })

            continue

        counts, classified_hits = calculate_taxonomic_support(
            qualifying_hits,
            taxonomy
        )

        category, reason = make_decision(
            qualifying_hits,
            counts
        )

        best = qualifying_hits[0]

        best_species = extract_species(best["stitle"])
        best_taxid = extract_taxid(best["stitle"])
        best_tax = taxonomy.get(best_taxid)

        if best_tax:
            best_lineage = (
                best_tax["superkingdom"] + "; "
                + best_tax["kingdom"] + "; "
                + best_tax["phylum"] + "; "
                + best_tax["class"] + "; "
                + best_tax["order"]
            )
        else:
            best_lineage = ""

        total = actual_qualifying_count

        arthropod_fraction = counts.get("ARTHROPOD", 0) / total
        bacterial_fraction = counts.get("BACTERIAL", 0) / total
        fungal_fraction = counts.get("FUNGAL", 0) / total
        non_arthropod_euk_fraction = counts.get(
            "NON_ARTHROPOD_EUKARYOTE", 0
        ) / total
        other_prokaryote_fraction = counts.get(
            "OTHER_PROKARYOTE", 0
        ) / total

        results.append({
            "gene_id": query_id,
            "best_hit_accession": best["sseqid"],
            "best_hit_species": best_species,
            "best_hit_taxid": best_taxid,
            "best_hit_lineage": best_lineage,
            "best_identity_pct": best["pident"],
            "best_query_coverage_pct": best["qcov"],
            "best_evalue": best["evalue"],
            "best_bitscore": best["bitscore"],
            "qualifying_hits": total,
            "total_blast_hits": total_blast_hits,
            "arthropod_hits": counts.get("ARTHROPOD", 0),
            "bacterial_hits": counts.get("BACTERIAL", 0),
            "fungal_hits": counts.get("FUNGAL", 0),
            "non_arthropod_eukaryote_hits": counts.get(
                "NON_ARTHROPOD_EUKARYOTE", 0
            ),
            "other_prokaryote_hits": counts.get(
                "OTHER_PROKARYOTE", 0
            ),
            "unknown_taxonomy_hits": counts.get("UNKNOWN_TAXONOMY", 0),
            "known_taxonomy_hits": total - counts.get(
                "UNKNOWN_TAXONOMY", 0
            ),
            "arthropod_fraction": arthropod_fraction,
            "bacterial_fraction": bacterial_fraction,
            "fungal_fraction": fungal_fraction,
            "non_arthropod_eukaryote_fraction": non_arthropod_euk_fraction,
            "other_prokaryote_fraction": other_prokaryote_fraction,
            "classification": category,
            "decision": category,
            "reason": reason,
            "best_hit_description": best["stitle"],
        })

    return results

def save_tsv(rows, output_path):
    if not rows:
        return

    fields = list(rows[0].keys())

    with open(output_path, "w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fields, delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)

    print(f"Saved TSV: {output_path}")

def save_excel(results, output_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("\nWARNING: openpyxl is not installed.")
        print("Install with:")
        print("pip install openpyxl")
        return

    if not results:
        print("No results to write to Excel.")
        return

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "All Results"

    fields = list(results[0].keys())

    for col, field in enumerate(fields, start=1):
        cell = sheet.cell(row=1, column=col, value=field)
        cell.font = Font(bold=True)

    for row_number, record in enumerate(results, start=2):
        for col, field in enumerate(fields, start=1):
            sheet.cell(row=row_number, column=col, value=record[field])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    categories = sorted(set(row["classification"] for row in results))

    for category in categories:
        ws = workbook.create_sheet(title=category[:31])

        category_rows = [
            row for row in results
            if row["classification"] == category
        ]

        for col, field in enumerate(fields, start=1):
            cell = ws.cell(row=1, column=col, value=field)
            cell.font = Font(bold=True)

        for row_number, record in enumerate(category_rows, start=2):
            for col, field in enumerate(fields, start=1):
                ws.cell(row=row_number, column=col, value=record[field])

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    workbook.save(output_path)
    print(f"Saved Excel: {output_path}")

def save_category_files(results, output_dir):
    categories = sorted(set(row["classification"] for row in results))

    for category in categories:
        category_rows = [row for row in results if row["classification"] == category]
        filename = category + ".tsv"
        output_path = os.path.join(output_dir, filename)
        save_tsv(category_rows, output_path)

def read_fasta(path):
    sequences = {}
    current_id = None
    sequence_parts = []

    with open(path, "r", encoding="utf-8", errors="replace") as file:
        for line in file:
            line = line.strip()
            if not line:
                continue

            if line.startswith(">"):
                if current_id is not None:
                    sequences[current_id] = "".join(sequence_parts)

                header = line[1:]
                current_id = header.split()[0]
                sequence_parts = []
            else:
                sequence_parts.append(line)

    if current_id is not None:
        sequences[current_id] = "".join(sequence_parts)

    return sequences

def write_fasta(sequence_ids, sequences, output_path):
    written = 0
    missing = []

    with open(output_path, "w", encoding="utf-8") as file:
        for sequence_id in sequence_ids:
            if sequence_id not in sequences:
                missing.append(sequence_id)
                continue

            file.write(">" + sequence_id + "\n")
            seq = sequences[sequence_id]

            for i in range(0, len(seq), 80):
                file.write(seq[i:i + 80] + "\n")

            written += 1

    print(f"\nFASTA sequences written: {written:,}")

    if missing:
        print(f"WARNING: {len(missing):,} IDs were not found in FASTA.")
        missing_file = output_path + ".missing.txt"

        with open(missing_file, "w", encoding="utf-8") as file:
            for sequence_id in missing:
                file.write(sequence_id + "\n")

def ask_categories():
    print("\n" + "=" * 70)
    print("FASTA CATEGORY SELECTION")
    print("=" * 70)

    for number, category in CATEGORIES.items():
        print(f"{number} = {category}")

    print("\nSelect one or multiple categories.")
    print("Example: 1")
    print("Example: 1,6")
    print("Example: 2,3")

    while True:
        answer = input("\nEnter category numbers: ").strip()

        if not answer:
            print("Please enter at least one category.")
            continue

        try:
            numbers = [int(x.strip()) for x in answer.split(",") if x.strip()]
        except ValueError:
            print("Invalid input. Use numbers such as 1,6.")
            continue

        invalid = [n for n in numbers if n not in CATEGORIES]

        if invalid:
            print("Invalid category number(s): " + ", ".join(str(x) for x in invalid))
            continue

        numbers = list(dict.fromkeys(numbers))
        return [CATEGORIES[n] for n in numbers]

def export_selected_fasta(results, fasta_sequences, categories, output_dir, output_prefix):
    selected_rows = [row for row in results if row["classification"] in categories]
    selected_ids = [row["gene_id"] for row in selected_rows]

    if not selected_ids:
        print("\nNo sequences belong to the selected categories.")
        return

    category_text = "_".join(categories)
    category_text = re.sub(r"[^A-Za-z0-9_.-]+", "_", category_text)

    output_path = os.path.join(output_dir, f"{output_prefix}_{category_text}filtered.fasta")
    write_fasta(selected_ids, fasta_sequences, output_path)

    print("\nSelected categories:")
    for category in categories:
        count = sum(1 for row in results if row["classification"] == category)
        print(f"  {category}: {count:,}")

    print("\nFASTA saved:")
    print(output_path)

def print_summary(results):
    counts = Counter(row["classification"] for row in results)

    print("\n" + "=" * 70)
    print("CONTAMINATION SCREEN SUMMARY")
    print("=" * 70)
    print(f"Proteins screened: {len(results):,}")

    for category in [
        "KEEP_ARTHROPOD",
        "POSSIBLE_BACTERIAL",
        "POSSIBLE_FUNGAL",
        "NON_ARTHROPOD_EUKARYOTE",
        "OTHER_PROKARYOTE",
        "AMBIGUOUS",
        "NO_QUALIFYING_HIT",
        "UNKNOWN_TAXONOMY",
    ]:
        count = counts.get(category, 0)
        percentage = count / len(results) * 100 if results else 0
        print(f"{category:30s} {count:8,} ({percentage:6.2f}%)")

def ask_output_name():
    print("\n" + "=" * 70)
    print("OUTPUT NAME")
    print("=" * 70)
    print("The output folder will be created beside the RAW BLAST file.")
    print("Enter a name for this screening run.")
    print("Example: metisa_screen_01")

    while True:
        name = input("\nOutput name: ").strip()
        name = re.sub(r'[<>:"/\\|?*]', "_", name).strip(". ")

        if not name:
            print("Please enter an output name.")
            continue

        return name

def main():
    print("\n" + "=" * 70)
    print("BLAST CONTAMINATION SCREEN")
    print("=" * 70)

    print("\nWORKFLOW:")
    print("1. Read RAW BLAST TSV")
    print("2. Count the actual BLAST hits for each protein")
    print("3. Apply identity / coverage / E-value thresholds")
    print("4. Use ALL qualifying hits actually present")
    print("5. Taxonomy screening ONLY on qualifying hits")
    print("6. Generate TSV + Excel + selectable FASTA")

    print("\nIMPORTANT:")
    print("The script does NOT assume a fixed number of BLAST hits.")
    print("It works whether each protein has 1, 5, 7, 10, 20, 40, "
          "or another number of hits.")

    print("\nTaxonomic categories:")
    print("  KEEP_ARTHROPOD")
    print("  POSSIBLE_BACTERIAL")
    print("  POSSIBLE_FUNGAL")
    print("  NON_ARTHROPOD_EUKARYOTE")
    print("  OTHER_PROKARYOTE")
    print("  AMBIGUOUS")
    print("  NO_QUALIFYING_HIT")
    print("  UNKNOWN_TAXONOMY")

    print("\nThresholds:")
    print(f"  Identity  >= {MIN_IDENTITY}%")
    print(f"  Coverage  >= {MIN_COVERAGE}%")
    print(f"  E-value   <= {MAX_EVALUE}")
    print(f"  Taxonomic dominance >= {DOMINANCE_THRESHOLD:.0%}")

    blast_file = expand_path(input("\nPath to RAW BLAST TSV: "))
    if not check_file_exists(blast_file, "RAW BLAST TSV"):
        return

    fasta_file = expand_path(input("\nPath to cleaned protein FASTA: "))
    if not check_file_exists(fasta_file, "Protein FASTA"):
        return

    rankedlineage_file = expand_path(
        input("\nPath to rankedlineage.dmp: ")
    )
    if not check_file_exists(rankedlineage_file, "rankedlineage.dmp"):
        return

    output_prefix = ask_output_name()
    blast_dir = os.path.dirname(blast_file)
    output_dir = os.path.join(blast_dir, output_prefix)

    if os.path.exists(output_dir):
        print("\nWARNING:")
        print("Output folder already exists:")
        print(output_dir)

        answer = input(
            "\nContinue and overwrite files in this folder? (y/n): "
        ).strip().lower()

        if answer != "y":
            print("Cancelled.")
            return

    os.makedirs(output_dir, exist_ok=True)

    raw_rows = read_blast_file(blast_file)

    if not raw_rows:
        print("\nERROR: No BLAST rows found.")
        return

    raw_grouped = group_hits_by_query(raw_rows)

    print("\nActual BLAST hit distribution:")

    hit_counts = Counter(len(hits) for hits in raw_grouped.values())

    for number_of_hits in sorted(hit_counts):
        protein_count = hit_counts[number_of_hits]
        print(f"  {number_of_hits:>4} hit(s): "
              f"{protein_count:,} protein(s)")

    qualifying_rows, failed_rows = apply_quality_filter(raw_rows)

    filtered_blast_path = os.path.join(
        output_dir,
        f"{output_prefix}_qualifying_hits.tsv"
    )
    save_tsv(qualifying_rows, filtered_blast_path)

    print("\nExtracting TaxIDs ONLY from quality-passing hits...")

    required_taxids = set()

    for row in qualifying_rows:
        taxid = extract_taxid(row["stitle"])
        if taxid is not None:
            required_taxids.add(taxid)

    print(f"Unique qualifying TaxIDs: {len(required_taxids):,}")

    taxonomy = load_required_taxonomy(
        rankedlineage_file,
        required_taxids
    )

    results = screen_proteins(raw_rows, taxonomy)

    all_tsv = os.path.join(
        output_dir,
        f"{output_prefix}_screening_all.tsv"
    )
    save_tsv(results, all_tsv)

    excel_file = os.path.join(
        output_dir,
        f"{output_prefix}_screening.xlsx"
    )
    save_excel(results, excel_file)

    print_summary(results)

    print("\nReading original protein FASTA...")

    fasta_sequences = read_fasta(fasta_file)

    print(f"FASTA sequences loaded: {len(fasta_sequences):,}")

    categories = ask_categories()

    export_selected_fasta(
        results,
        fasta_sequences,
        categories,
        output_dir,
        output_prefix
    )

    print("\n" + "=" * 70)
    print("SCREENING COMPLETE")
    print("=" * 70)

    print("\nOutput folder:")
    print(output_dir)

    print("\nImportant:")
    print("The number of BLAST hits is determined from the actual "
          "input file.")
    print("No fixed hit count such as 5, 7, 10, 20, or 40 is assumed.")
    print("Only hits passing identity + coverage + E-value "
          "are used for taxonomy.")
    print("Unknown-taxonomy hits are included from the denominator "
          "when calculating taxonomic dominance.")
    print("\nThe selected FASTA can be used for "
          "downstream annotation such as InterProScan.")

if __name__ == "__main__":
    main()