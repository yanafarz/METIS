import csv
import math
import os
import re
from collections import defaultdict

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

DEFAULT_MAX_CANDIDATES = 500

MIN_IDENTITY_MINING = 30.0
MIN_QUERY_COVERAGE_MINING = 70.0
MAX_EVALUE_MINING = 1e-10

STRONG_IDENTITY_MINING = 40.0
STRONG_COVERAGE_MINING = 80.0
STRONG_EVALUE_MINING = 1e-20

MIN_ARTHROPOD_FRACTION_KEEP = 0.70
MIN_ARTHROPOD_FRACTION_STRONG = 0.90

MAX_HITS_PER_GENE = 50

WEIGHT_HOMOLOGY = 30.0
WEIGHT_TAXONOMY = 25.0
WEIGHT_ANNOTATION = 15.0
WEIGHT_INTERPRO = 15.0
WEIGHT_TARGET_CLASS = 15.0

CANDIDATE_COLUMNS = [
    "gene_id", "rank", "final_score", "decision", "reason",
    "sequence_length", "best_hit_accession", "best_hit_species",
    "best_hit_taxid", "best_hit_lineage", "best_identity_pct",
    "best_query_coverage_pct", "best_evalue", "best_bitscore",
    "qualifying_hits", "total_blast_hits", "arthropod_fraction",
    "bacterial_fraction", "fungal_fraction",
    "non_arthropod_eukaryote_fraction", "other_prokaryote_fraction",
    "taxonomy_classification", "homology_score", "taxonomy_score",
    "annotation_score", "interpro_score", "target_class_score",
    "functional_description", "interpro_accessions",
    "interpro_descriptions", "pfam", "panther", "go_terms",
    "functional_keywords", "target_class", "best_hit_description",
]

EVIDENCE_COLUMNS = [
    "gene_id", "hit_rank", "sseqid", "pident", "qcov", "evalue",
    "bitscore", "length", "qlen", "slen", "qualifies", "strong_hit",
    "stitle",
]

REJECTED_COLUMNS = [
    x for x in CANDIDATE_COLUMNS
    if x != "rank"
]

def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def safe_float(value, default=0.0):
    try:
        return float(str(value).strip())
    except Exception:
        return default


def safe_int(value, default=0):
    try:
        return int(float(str(value).strip()))
    except Exception:
        return default


def normalize_gene_id(value):
    value = clean_text(value)

    if value.startswith(">"):
        value = value[1:]

    if not value:
        return ""

    return value.split()[0]


def normalize_column_name(name):
    name = clean_text(name).lower()
    name = name.replace(" ", "")
    name = name.replace("-", "_")
    name = name.replace(".", "_")
    return name


def find_column(fieldnames, candidates):
    if not fieldnames:
        return None

    normalized = {
        normalize_column_name(x): x
        for x in fieldnames
    }

    for candidate in candidates:
        key = normalize_column_name(candidate)

        if key in normalized:
            return normalized[key]

    return None


def safe_filename(name):
    name = clean_text(name)
    name = os.path.basename(name)

    for suffix in [
        ".tsv",
        ".txt",
        ".xlsx",
        ".csv"
    ]:
        if name.lower().endswith(suffix):
            name = name[:-len(suffix)]

    if not name:
        name = "metisa_results"

    return name


def ask_existing_file(prompt):
    while True:
        path = input(prompt).strip().strip('"')

        if not path:
            print("Please enter a file path.")
            continue

        path = os.path.abspath(
            os.path.expanduser(path)
        )

        if os.path.isfile(path):
            return path

        print("File not found:")
        print(path)
        print()


def ask_output_basename():
    while True:
        value = input(
            "Output base name [metisa_results]: "
        ).strip().strip('"')

        if not value:
            value = "metisa_results"

        value = safe_filename(value)

        if value:
            return value

        print("Please enter a valid filename.")

def ask_max_candidates():
    while True:
        value = input(
            f"Number of final candidates [{DEFAULT_MAX_CANDIDATES}]: "
        ).strip()

        if not value:
            return DEFAULT_MAX_CANDIDATES

        try:
            value = int(value)

            if value <= 0:
                print("Please enter a number greater than 0.")
                continue

            return value

        except ValueError:
            print("Please enter a whole number.")

def print_header(title):
    print()
    print("=" * 78)
    print(title)
    print("=" * 78)


def join_unique(values, separator="; "):
    seen = set()
    result = []

    for value in values:
        value = clean_text(value)

        if not value:
            continue

        if value in seen:
            continue

        seen.add(value)
        result.append(value)

    return separator.join(result)


def truncate_text(value, maximum=1000):
    value = clean_text(value)

    if len(value) <= maximum:
        return value

    return value[:maximum - 3] + "..."


def write_tsv(path, rows, fieldnames):
    with open(
        path,
        "w",
        encoding="utf-8",
        newline=""
    ) as handle:

        writer = csv.DictWriter(
            handle,
            fieldnames=fieldnames,
            delimiter="\t",
            extrasaction="ignore"
        )

        writer.writeheader()

        for row in rows:
            writer.writerow(row)


def write_text(path, text):
    with open(
        path,
        "w",
        encoding="utf-8"
    ) as handle:
        handle.write(text)

def read_fasta(path):
    sequences = {}

    current_id = None
    chunks = []

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="replace"
    ) as handle:

        for raw_line in handle:

            line = raw_line.strip()

            if not line:
                continue

            if line.startswith(">"):

                if current_id is not None:
                    sequences[current_id] = "".join(chunks)

                current_id = normalize_gene_id(line)
                chunks = []

            else:
                chunks.append(
                    line.replace(" ", "")
                )

    if current_id is not None:
        sequences[current_id] = "".join(chunks)

    return sequences


def validate_sequences(sequences):

    valid_amino_acids = set(
        "ACDEFGHIKLMNPQRSTVWY"
        "BJOUXZ"
    )

    stats = {
        "total": len(sequences),
        "empty": 0,
        "internal_stop": 0,
        "invalid_characters": 0,
    }

    for sequence in sequences.values():

        sequence = sequence.upper()

        if not sequence:
            stats["empty"] += 1
            continue

        if "*" in sequence[:-1]:
            stats["internal_stop"] += 1

        invalid = (
            set(sequence)
            - valid_amino_acids
            - {"*"}
        )

        if invalid:
            stats["invalid_characters"] += 1

    return stats

def parse_blast(path):
    hits = defaultdict(list)

    with open(
        path, "r", encoding="utf-8-sig",
        errors="replace", newline=""
    ) as handle:

        reader = csv.DictReader(
            handle, delimiter="\t"
        )

        if not reader.fieldnames:
            raise ValueError(
                "BLAST file does not contain a header."
            )

        fields = reader.fieldnames

        qseqid_col = find_column(
            fields,
            ["qseqid", "query", "query_id", "gene_id"]
        )

        sseqid_col = find_column(
            fields,
            ["sseqid", "subject", "subject_id", "accession"]
        )

        pident_col = find_column(
            fields,
            ["pident", "identity", "identity_pct"]
        )

        length_col = find_column(
            fields,
            ["length", "alignment_length"]
        )

        qlen_col = find_column(
            fields,
            ["qlen", "query_length"]
        )

        slen_col = find_column(
            fields,
            ["slen", "subject_length"]
        )

        qcov_col = find_column(
            fields,
            ["qcov", "query_coverage", "best_query_coverage"]
        )

        evalue_col = find_column(
            fields,
            ["evalue", "e_value"]
        )

        bitscore_col = find_column(
            fields,
            ["bitscore", "bit_score"]
        )

        title_col = find_column(
            fields,
            ["stitle", "description",
             "subject_title", "best_hit_description"]
        )

        required = {
            "qseqid": qseqid_col,
            "sseqid": sseqid_col,
            "pident": pident_col,
            "evalue": evalue_col,
            "bitscore": bitscore_col,
        }

        missing = [
            name
            for name, column in required.items()
            if column is None
        ]

        if missing:
            raise ValueError(
                "BLAST file is missing required columns: "
                + ", ".join(missing)
            )

        total_rows = 0

        for row in reader:
            total_rows += 1

            gene_id = normalize_gene_id(
                row.get(qseqid_col)
            )

            if not gene_id:
                continue

            hit = {
                "qseqid": gene_id,

                "sseqid": clean_text(
                    row.get(sseqid_col)
                ),

                "pident": safe_float(
                    row.get(pident_col)
                ),

                "length": (
                    safe_int(row.get(length_col))
                    if length_col else 0
                ),

                "qlen": (
                    safe_int(row.get(qlen_col))
                    if qlen_col else 0
                ),

                "slen": (
                    safe_int(row.get(slen_col))
                    if slen_col else 0
                ),

                "qcov": (
                    safe_float(row.get(qcov_col))
                    if qcov_col else 0.0
                ),

                "evalue": safe_float(
                    row.get(evalue_col)
                ),

                "bitscore": safe_float(
                    row.get(bitscore_col)
                ),

                "stitle": (
                    clean_text(row.get(title_col))
                    if title_col else ""
                ),
            }

            hits[gene_id].append(hit)

    return hits, total_rows

def qualify_blast_hit(hit):
    return (
        hit["pident"] >= MIN_IDENTITY_MINING
        and
        hit["qcov"] >= MIN_QUERY_COVERAGE_MINING
        and
        hit["evalue"] <= MAX_EVALUE_MINING
    )


def strong_blast_hit(hit):
    return (
        hit["pident"] >= STRONG_IDENTITY_MINING
        and
        hit["qcov"] >= STRONG_COVERAGE_MINING
        and
        hit["evalue"] <= STRONG_EVALUE_MINING
    )

def blast_hit_strength(hit):
    identity = hit["pident"]
    coverage = hit["qcov"]
    evalue = hit["evalue"]

    identity_score = (
        (identity - MIN_IDENTITY_MINING)
        / (70.0 - MIN_IDENTITY_MINING)
        * 100.0
    )

    identity_score = max(
        0.0,
        min(100.0, identity_score)
    )

    coverage_score = (
        (coverage - MIN_QUERY_COVERAGE_MINING)
        / (100.0 - MIN_QUERY_COVERAGE_MINING)
        * 100.0
    )

    coverage_score = max(
        0.0,
        min(100.0, coverage_score)
    )

    if evalue <= 0:
        evalue_score = 100.0
    else:
        neglog = -math.log10(
            max(evalue, 1e-300)
        )

        # 1e-10 -> ~33
        # 1e-20 -> ~67
        # 1e-30 -> 100
        evalue_score = (
            neglog
            / 30.0
            * 100.0
        )

        evalue_score = max(
            0.0,
            min(100.0, evalue_score)
        )

    return (
        identity_score * 0.50
        + coverage_score * 0.30
        + evalue_score * 0.20
    )


def select_best_hit(hits):
    qualifying = [
        hit
        for hit in hits
        if qualify_blast_hit(hit)
    ]

    if not qualifying:
        return None, []

    qualifying.sort(
        key=lambda x: (
            x["bitscore"],
            x["pident"],
            x["qcov"],
            -math.log10(
                max(x["evalue"], 1e-300)
            )
        ),
        reverse=True
    )

    return qualifying[0], qualifying

def parse_interpro(path):
    annotations = defaultdict(
        lambda: {
            "interpro_accessions": [],
            "interpro_descriptions": [],
            "signature_accessions": [],
            "signature_descriptions": [],
            "pfam": [],
            "panther": [],
            "go_terms": [],
            "pathways": [],
        }
    )

    with open(
        path,
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline=""
    ) as handle:

        reader = csv.DictReader(
            handle,
            delimiter="\t"
        )

        if not reader.fieldnames:
            raise ValueError(
                "InterPro file does not contain a header."
            )

        fields = reader.fieldnames

        protein_col = find_column(
            fields,
            [
                "Protein Accession",
                "protein_accession",
                "protein",
                "accession"
            ]
        )

        analysis_col = find_column(
            fields,
            [
                "Analysis",
                "analysis"
            ]
        )

        signature_col = find_column(
            fields,
            [
                "Signature accession",
                "signature_accession"
            ]
        )

        signature_desc_col = find_column(
            fields,
            [
                "Signature description",
                "signature_description"
            ]
        )

        start_col = find_column(
            fields,
            [
                "Start location",
                "start"
            ]
        )

        stop_col = find_column(
            fields,
            [
                "Stop location",
                "stop"
            ]
        )

        score_col = find_column(
            fields,
            [
                "Score",
                "score"
            ]
        )

        interpro_col = find_column(
            fields,
            [
                "InterPro accession",
                "interpro_accession"
            ]
        )

        interpro_desc_col = find_column(
            fields,
            [
                "InterPro description",
                "interpro_description"
            ]
        )

        go_col = find_column(
            fields,
            [
                "GO terms",
                "go_terms",
                "go"
            ]
        )

        pathway_col = find_column(
            fields,
            [
                "Pathways",
                "pathways",
                "pathway"
            ]
        )

        if protein_col is None:
            raise ValueError(
                "InterPro file is missing "
                "'Protein Accession'."
            )

        for row in reader:
            gene_id = normalize_gene_id(
                row.get(protein_col)
            )

            if not gene_id:
                continue

            annotation = annotations[gene_id]

            analysis = (
                clean_text(row.get(analysis_col))
                if analysis_col else ""
            )

            signature = (
                clean_text(row.get(signature_col))
                if signature_col else ""
            )

            signature_desc = (
                clean_text(
                    row.get(signature_desc_col)
                )
                if signature_desc_col else ""
            )

            interpro = (
                clean_text(row.get(interpro_col))
                if interpro_col else ""
            )

            interpro_desc = (
                clean_text(
                    row.get(interpro_desc_col)
                )
                if interpro_desc_col else ""
            )

            go = (
                clean_text(row.get(go_col))
                if go_col else ""
            )

            pathway = (
                clean_text(row.get(pathway_col))
                if pathway_col else ""
            )

            if interpro and interpro != "-":
                annotation[
                    "interpro_accessions"
                ].append(interpro)

            if interpro_desc and interpro_desc != "-":
                annotation[
                    "interpro_descriptions"
                ].append(interpro_desc)

            if signature and signature != "-":
                annotation[
                    "signature_accessions"
                ].append(signature)

            if signature_desc and signature_desc != "-":
                annotation[
                    "signature_descriptions"
                ].append(signature_desc)

            if analysis.upper() == "PFAM":
                if signature and signature != "-":
                    annotation["pfam"].append(
                        signature
                    )

                if signature_desc and signature_desc != "-":
                    annotation["pfam"].append(
                        signature_desc
                    )

            if analysis.upper() == "PANTHER":
                if signature and signature != "-":
                    annotation["panther"].append(
                        signature
                    )

                if signature_desc and signature_desc != "-":
                    annotation["panther"].append(
                        signature_desc
                    )

            if go and go != "-":
                annotation["go_terms"].append(go)

            if pathway and pathway != "-":
                annotation["pathways"].append(
                    pathway
                )

    return annotations

def has_real_interpro_evidence(annotation):

    if not annotation:
        return False

    evidence_keys = [
        "interpro_accessions",
        "interpro_descriptions",
        "signature_accessions",
        "signature_descriptions",
        "pfam",
        "panther",
        "go_terms",
        "pathways",
    ]

    for key in evidence_keys:
        values = annotation.get(key, [])

        for value in values:
            value = clean_text(value)

            if value and value != "-":
                return True

    return False

def parse_taxonomy(path):

    taxonomy = {}

    with open(
        path,
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline=""
    ) as handle:

        reader = csv.DictReader(
            handle,
            delimiter="\t"
        )

        if not reader.fieldnames:
            raise ValueError(
                "Taxonomy file does not contain a header."
            )

        fields = reader.fieldnames

        gene_col = find_column(
            fields,
            [
                "gene_id",
                "qseqid",
                "protein_accession"
            ]
        )

        if gene_col is None:
            raise ValueError(
                "Taxonomy file is missing gene_id."
            )

        for row in reader:

            gene_id = normalize_gene_id(
                row.get(gene_col)
            )

            if not gene_id:
                continue

            item = {}

            for field in fields:
                item[
                    normalize_column_name(field)
                ] = clean_text(
                    row.get(field)
                )

            taxonomy[gene_id] = item

    return taxonomy


def taxonomy_value(item, names, default=""):

    if not item:
        return default

    for name in names:

        key = normalize_column_name(name)

        if key in item:
            return item[key]

    return default


def taxonomy_fraction(item, names):

    return safe_float(
        taxonomy_value(
            item,
            names,
            "0"
        ),
        0.0
    )

def classify_taxonomy(
    item,
    blast_qualifying_hits
):
    if item:
        arth_frac = taxonomy_fraction(
            item, ["arthropod_fraction"]
        )

        bacteria_frac = taxonomy_fraction(
            item,
            [
                "bacterial_fraction",
                "bacteria_fraction"
            ]
        )

        fungal_frac = taxonomy_fraction(
            item,
            [
                "fungal_fraction",
                "fungi_fraction"
            ]
        )

        nonarth_frac = taxonomy_fraction(
            item,
            ["non_arthropod_eukaryote_fraction"]
        )

        other_prok_frac = taxonomy_fraction(
            item,
            ["other_prokaryote_fraction"]
        )

        classification = taxonomy_value(
            item,
            [
                "taxonomy_classification",
                "classification"
            ]
        )

        decision = taxonomy_value(
            item, ["decision"]
        )

        return {
            "arthropod_fraction": arth_frac,
            "bacterial_fraction": bacteria_frac,
            "fungal_fraction": fungal_frac,
            "non_arthropod_eukaryote_fraction": nonarth_frac,
            "other_prokaryote_fraction": other_prok_frac,
            "classification": classification,
            "decision": decision,
        }

    arthropod = 0
    bacterial = 0
    fungal = 0
    nonarth = 0
    other_prok = 0
    unknown = 0

    for hit in blast_qualifying_hits:
        title = hit["stitle"].lower()

        if any(
            word in title
            for word in [
                "bacter", "escherichia", "bacillus",
                "streptococcus", "staphylococcus",
                "pseudomonas", "mycobacterium"
            ]
        ):
            bacterial += 1

        elif any(
            word in title
            for word in [
                "fung", "candida", "aspergillus",
                "saccharomyces", "yeast"
            ]
        ):
            fungal += 1

        elif any(
            word in title
            for word in [
                "insect", "lepidoptera", "diptera",
                "coleoptera", "hemiptera", "hymenoptera",
                "blattodea", "orthoptera", "arachnid",
                "arthropod"
            ]
        ):
            arthropod += 1

        elif any(
            word in title
            for word in [
                "plant", "arabidopsis", "rice",
                "maize", "soybean", "human", "mouse",
                "fish", "vertebrate", "mammal"
            ]
        ):
            nonarth += 1

        else:
            unknown += 1

    total = (
        arthropod
        + bacterial
        + fungal
        + nonarth
        + other_prok
        + unknown
    )

    if total == 0:
        total = 1

    arth_frac = arthropod / total
    bact_frac = bacterial / total
    fungal_frac = fungal / total
    nonarth_frac = nonarth / total
    other_frac = other_prok / total

    if arth_frac >= MIN_ARTHROPOD_FRACTION_KEEP:
        classification = "KEEP_ARTHROPOD"
    elif bact_frac >= 0.50:
        classification = "LIKELY_BACTERIAL"
    elif fungal_frac >= 0.50:
        classification = "LIKELY_FUNGAL"
    elif nonarth_frac >= 0.50:
        classification = "NON_ARTHROPOD_EUKARYOTE"
    else:
        classification = "MIXED_OR_UNKNOWN"

    return {
        "arthropod_fraction": arth_frac,
        "bacterial_fraction": bact_frac,
        "fungal_fraction": fungal_frac,
        "non_arthropod_eukaryote_fraction": nonarth_frac,
        "other_prokaryote_fraction": other_frac,
        "classification": classification,
        "decision": classification,
    }

def annotation_quality(annotation):

    if not annotation:
        return 0.0

    score = 0.0

    if annotation["interpro_accessions"]:
        score += 35.0

    if annotation["interpro_descriptions"]:
        score += 20.0

    if annotation["pfam"]:
        score += 20.0

    if annotation["panther"]:
        score += 15.0

    if annotation["go_terms"]:
        score += 7.0

    if annotation["pathways"]:
        score += 3.0

    return min(100.0, score)

def annotation_keywords(annotation):
    if not annotation:
        return []

    text_parts = []

    for key in [
        "interpro_descriptions",
        "signature_descriptions",
        "pfam",
        "panther"
    ]:
        text_parts.extend(
            annotation.get(key, [])
        )

    text = " ".join(text_parts).lower()

    groups = {
        "enzyme": [
            "enzyme",
            "hydrolase",
            "transferase",
            "oxidoreductase",
            "protease",
            "peptidase",
            "phosphatase",
            "kinase"
        ],

        "transcription_factor": [
            "transcription factor",
            "transcriptional regulator",
            "dna-binding",
            "myb",
            "madf",
            "zinc finger"
        ],

        "receptor": [
            "receptor",
            "ligand-binding"
        ],

        "transporter": [
            "transporter",
            "transport",
            "channel"
        ],

        "membrane": [
            "transmembrane",
            "membrane protein"
        ],

        "development": [
            "development",
            "developmental",
            "growth",
            "cell differentiation"
        ],

        "digestion": [
            "digest",
            "amylase",
            "lipase",
            "trypsin",
            "chymotrypsin",
            "protease",
            "peptidase",
            "carbohydrase",
            "glycosidase"
        ],

        "hormone": [
            "hormone",
            "ecdysone",
            "juvenile hormone"
        ],

        "signal": [
            "signal peptide",
            "signaling",
            "signalling"
        ],
    }

    result = []

    for group, terms in groups.items():
        if any(
            term in text
            for term in terms
        ):
            result.append(group)

    return result

def calculate_homology_score(
    best_hit,
    qualifying_hits
):

    if best_hit is None:
        return 0.0

    best_strength = blast_hit_strength(
        best_hit
    )

    n = len(qualifying_hits)

    hit_support = (
        1.0
        - math.exp(-n / 3.0)
    )

    hit_count_score = (
        hit_support * 100.0
    )

    strong_hits = sum(
        1
        for hit in qualifying_hits
        if strong_blast_hit(hit)
    )

    strong_support = (
        1.0
        - math.exp(-strong_hits / 2.0)
    )

    strong_score = (
        strong_support * 100.0
    )

    # Best hit remains dominant.
    score = (
        best_strength * 0.70
        +
        hit_count_score * 0.15
        +
        strong_score * 0.15
    )

    return min(
        100.0,
        score
    )

def classify_target_class(
    best_hit,
    annotation
):

    text_parts = []

    # BLAST description
    if best_hit:
        text_parts.append(
            clean_text(
                best_hit.get("stitle", "")
            )
        )

    # InterPro / Pfam / Panther descriptions
    if annotation:
        for key in [
            "interpro_descriptions",
            "signature_descriptions",
            "pfam",
            "panther",
        ]:
            text_parts.extend(
                annotation.get(key, [])
            )

    text = " ".join(
        clean_text(x)
        for x in text_parts
    ).lower()

    chitin_terms = [
        "chitin synthase",
        "chitinase",
        "chitin binding",
        "chitin-binding",
        "chitin metabolism",
        "chitin biosynthesis",
        "chitin degradation",
        "chitinase activity",
        "chitin synthase activity",
        "chitin",
    ]

    if any(
        term in text
        for term in chitin_terms
    ):
        return "CHITIN"

    development_terms = [
        "ecdysone",
        "ecdysteroid",
        "juvenile hormone",
        "juvenile hormone binding",
        "juvenile hormone esterase",
        "juvenile hormone receptor",
        "development",
        "developmental",
        "growth",
        "molting",
        "moulting",
        "metamorphosis",
        "cell differentiation",
    ]

    if any(
        term in text
        for term in development_terms
    ):
        return "GROWTH_DEVELOPMENT"

    digestion_terms = [
        "digestive enzyme",
        "digestion",
        "amylase",
        "alpha-amylase",
        "lipase",
        "trypsin",
        "trypsin-like",
        "chymotrypsin",
        "chymotrypsin-like",
        "peptidase",
        "protease",
        "carboxypeptidase",
        "aminopeptidase",
        "glycosidase",
        "carbohydrase",
        "cellulase",
    ]

    if any(
        term in text
        for term in digestion_terms
    ):
        return "DIGESTION"

    detox_terms = [
        "cytochrome p450",
        "cytochrome p450 monooxygenase",
        "glutathione s-transferase",
        "glutathione transferase",
        "gst",
        "carboxylesterase",
        "esterase",
        "detoxification",
        "detoxification enzyme",
        "xenobiotic",
        "xenobiotic metabolism",
        "udp-glucuronosyltransferase",
        "udp-glycosyltransferase",
        "ug t",
        "aldo-keto reductase",
    ]

    if any(
        term in text
        for term in detox_terms
    ):
        return "DETOXIFICATION"

    signaling_terms = [
        "receptor",
        "ligand-binding receptor",
        "g-protein coupled receptor",
        "gpcr",
        "signaling",
        "signalling",
        "signal transduction",
        "protein kinase",
        "kinase",
    ]

    if any(
        term in text
        for term in signaling_terms
    ):
        return "SIGNALING"

    transport_terms = [
        "transporter",
        "transport protein",
        "ion channel",
        "channel protein",
        "abc transporter",
        "atp-binding cassette",
        "membrane transporter",
    ]

    if any(
        term in text
        for term in transport_terms
    ):
        return "TRANSPORT"

    regulation_terms = [
        "transcription factor",
        "transcriptional regulator",
        "dna-binding",
        "transcriptional activator",
        "transcriptional repressor",
        "zinc finger",
    ]

    if any(
        term in text
        for term in regulation_terms
    ):
        return "TRANSCRIPTION_REGULATION"

    return "OTHER"


TARGET_CLASS_SCORES = {
    "CHITIN": 100.0,
    "GROWTH_DEVELOPMENT": 100.0,
    "DIGESTION": 95.0,
    "SIGNALING": 85.0,
    "DETOXIFICATION": 80.0,
    "TRANSPORT": 70.0,
    "TRANSCRIPTION_REGULATION": 65.0,
    "OTHER": 30.0,
}

def calculate_target_class_score(target_class):
    return TARGET_CLASS_SCORES.get(
        target_class,
        30.0
    )

def calculate_taxonomy_score(tax):

    if not tax:
        return 0.0

    arth = tax[
        "arthropod_fraction"
    ]

    bact = tax[
        "bacterial_fraction"
    ]

    fungal = tax[
        "fungal_fraction"
    ]

    nonarth = tax[
        "non_arthropod_eukaryote_fraction"
    ]

    other_prok = tax[
        "other_prokaryote_fraction"
    ]

    # Arthropod support.
    score = arth * 100.0

    # Strong contamination penalties.
    score -= bact * 100.0
    score -= fungal * 90.0
    score -= other_prok * 90.0

    # Non-arthropod eukaryotes are less alarming
    # than bacterial/fungal contamination.
    score -= nonarth * 45.0

    return max(
        0.0,
        min(100.0, score)
    )

def calculate_annotation_score(
    best_hit,
    annotation
):

    score = 0.0

    if best_hit:

        description = best_hit[
            "stitle"
        ].lower()

        if description:

            if "hypothetical protein" in description:
                score += 3.0

            elif "uncharacterized" in description:
                score += 5.0

            elif any(
                term in description
                for term in [
                    "trypsin",
                    "chymotrypsin",
                    "peptidase",
                    "protease",
                    "amylase",
                    "lipase",
                    "chitinase",
                    "chitin synthase",
                    "cytochrome p450",
                    "glutathione s-transferase",
                    "ecdysone",
                    "ecdysteroid",
                    "juvenile hormone",
                    "receptor",
                    "transporter",
                    "kinase",
                    "transcription factor",
                    "enzyme"
                ]
            ):
                score += 30.0

            else:
                score += 10.0

    return min(
        100.0,
        score
    )

def calculate_final_score(
    homology_score,
    taxonomy_score,
    annotation_score,
    interpro_score,
    target_class_score,
    tax
):

    score = (
        homology_score
        * WEIGHT_HOMOLOGY
        / 100.0
        +
        taxonomy_score
        * WEIGHT_TAXONOMY
        / 100.0
        +
        annotation_score
        * WEIGHT_ANNOTATION
        / 100.0
        +
        interpro_score
        * WEIGHT_INTERPRO
        / 100.0
        +
        target_class_score
        * WEIGHT_TARGET_CLASS
        / 100.0
    )

    # Additional contamination penalty.
    contamination = (
        tax["bacterial_fraction"] * 30.0
        +
        tax["fungal_fraction"] * 25.0
        +
        tax["other_prokaryote_fraction"] * 30.0
        +
        tax["non_arthropod_eukaryote_fraction"] * 10.0
    )

    score -= contamination

    return max(
        0.0,
        min(100.0, score)
    )

def determine_decision(
    best_hit,
    qualifying_hits,
    tax,
    annotation,
    final_score
):

    if best_hit is None:
        return (
            "REJECT",
            "No BLAST hit passed the identity, "
            "coverage and E-value filter."
        )

    arth = tax[
        "arthropod_fraction"
    ]

    bact = tax[
        "bacterial_fraction"
    ]

    fungal = tax[
        "fungal_fraction"
    ]

    nonarth = tax[
        "non_arthropod_eukaryote_fraction"
    ]

    classification = tax.get(
        "classification",
        ""
    ).upper()

    if bact >= 0.70:
        return (
            "REJECT",
            "Strong bacterial support among "
            "qualifying hits."
        )

    if fungal >= 0.70:
        return (
            "REJECT",
            "Strong fungal support among "
            "qualifying hits."
        )

    if classification in {
        "LIKELY_BACTERIAL",
        "LIKELY_FUNGAL"
    }:
        return (
            "REJECT",
            "Taxonomy classification indicates "
            "likely microbial contamination."
        )

    if (
        arth >= MIN_ARTHROPOD_FRACTION_STRONG
        and
        final_score >= 70
        and
        len(qualifying_hits) >= 2
    ):
        return (
            "KEEP",
            "Strong arthropod homology with multiple "
            "qualifying hits and strong combined evidence."
        )

    if (
        arth >= MIN_ARTHROPOD_FRACTION_KEEP
        and
        final_score >= 60
    ):
        return (
            "KEEP",
            "Good arthropod support and sufficient "
            "combined evidence."
        )

    if nonarth >= 0.60:
        return (
            "REVIEW",
            "Most qualifying hits are non-arthropod "
            "eukaryotes."
        )

    if arth >= 0.40:
        return (
            "REVIEW",
            "Arthropod support exists but taxonomy "
            "is mixed."
        )

    return (
        "REVIEW",
        "Evidence is insufficient for confident "
        "arthropod classification."
    )

def choose_functional_description(
    best_hit,
    annotation
):

    if annotation:

        descriptions = annotation.get(
            "interpro_descriptions",
            []
        )

        for description in descriptions:

            description = clean_text(
                description
            )

            if (
                description
                and
                description != "-"
            ):
                return description

        descriptions = annotation.get(
            "signature_descriptions",
            []
        )

        for description in descriptions:

            description = clean_text(
                description
            )

            if (
                description
                and
                description != "-"
            ):
                return description

    if best_hit:
        return best_hit["stitle"]

    return ""

def extract_species(title):

    title = clean_text(title)

    if not title:
        return ""

    # UniProt:
    # Protein name OS=Species OX=TaxID
    match = re.search(
        r"\bOS=([^=]+?)(?:\s+OX=|\s+GN=|\s+PE=|\s+SV=|$)",
        title
    )

    if match:
        return match.group(1).strip()

    # Generic fallback.
    match = re.search(
        r"\bOS=([A-Z][^=]+)",
        title
    )

    if match:
        return match.group(1).strip()

    return ""

def build_candidate(
    gene_id,
    blast_hits,
    annotation,
    taxonomy,
    sequences
):

    all_hits = sorted(
        blast_hits,
        key=lambda x: (
            x["bitscore"],
            x["pident"],
            x["qcov"],
            -math.log10(max(x["evalue"], 1e-300))
        ),
        reverse=True
    )[:MAX_HITS_PER_GENE]

    best_hit, qualifying_hits = select_best_hit(all_hits)

    tax_item = taxonomy.get(gene_id, {})
    tax = classify_taxonomy(tax_item, qualifying_hits)

    homology_score = calculate_homology_score(
        best_hit,
        qualifying_hits
    )

    taxonomy_score = calculate_taxonomy_score(tax)
    interpro_score = annotation_quality(annotation)

    annotation_score = calculate_annotation_score(
        best_hit,
        annotation
    )

    target_class = classify_target_class(
        best_hit,
        annotation
    )

    target_class_score = calculate_target_class_score(
        target_class
    )

    final_score = calculate_final_score(
        homology_score,
        taxonomy_score,
        annotation_score,
        interpro_score,
        target_class_score,
        tax
    )

    decision, reason = determine_decision(
        best_hit,
        qualifying_hits,
        tax,
        annotation,
        final_score
    )

    sequence = sequences.get(gene_id, "")
    description = choose_functional_description(
        best_hit,
        annotation
    )

    # --------------------------------------------------------
    # BEST HIT DETAILS
    # --------------------------------------------------------

    if best_hit:

        best_accession = best_hit["sseqid"]

        best_species = taxonomy_value(
            tax_item,
            ["best_hit_species"]
        ) or extract_species(
            best_hit["stitle"]
        )

        best_taxid = taxonomy_value(
            tax_item,
            ["best_hit_taxid"]
        )

        best_lineage = taxonomy_value(
            tax_item,
            ["best_hit_lineage"]
        )

        best_identity = best_hit["pident"]
        best_coverage = best_hit["qcov"]
        best_evalue = best_hit["evalue"]
        best_bitscore = best_hit["bitscore"]
        best_description = best_hit["stitle"]

    else:

        best_accession = ""
        best_species = ""
        best_taxid = ""
        best_lineage = ""
        best_identity = ""
        best_coverage = ""
        best_evalue = ""
        best_bitscore = ""
        best_description = ""

    row = {
        "gene_id": gene_id,
        "rank": "",
        "final_score": round(final_score, 3),
        "decision": decision,
        "reason": reason,
        "sequence_length": len(sequence),

        "best_hit_accession": best_accession,
        "best_hit_species": best_species,
        "best_hit_taxid": best_taxid,
        "best_hit_lineage": best_lineage,

        "best_identity_pct": best_identity,
        "best_query_coverage_pct": best_coverage,
        "best_evalue": best_evalue,
        "best_bitscore": best_bitscore,

        "qualifying_hits": len(qualifying_hits),
        "total_blast_hits": len(blast_hits),

        "arthropod_fraction": round(
            tax["arthropod_fraction"], 4
        ),
        "bacterial_fraction": round(
            tax["bacterial_fraction"], 4
        ),
        "fungal_fraction": round(
            tax["fungal_fraction"], 4
        ),
        "non_arthropod_eukaryote_fraction": round(
            tax["non_arthropod_eukaryote_fraction"], 4
        ),
        "other_prokaryote_fraction": round(
            tax["other_prokaryote_fraction"], 4
        ),

        "taxonomy_classification": tax.get(
            "classification",
            ""
        ),

        "homology_score": round(
            homology_score, 3
        ),
        "taxonomy_score": round(
            taxonomy_score, 3
        ),
        "annotation_score": round(
            annotation_score, 3
        ),
        "interpro_score": round(
            interpro_score, 3
        ),

        "functional_description": description,

        "interpro_accessions": join_unique(
            annotation.get(
                "interpro_accessions",
                []
            )
        ),
        "interpro_descriptions": join_unique(
            annotation.get(
                "interpro_descriptions",
                []
            )
        ),
        "pfam": join_unique(
            annotation.get("pfam", [])
        ),
        "panther": join_unique(
            annotation.get("panther", [])
        ),
        "go_terms": join_unique(
            annotation.get("go_terms", [])
        ),

        "functional_keywords": "; ".join(
            annotation_keywords(annotation)
        ),

        "target_class_score": round(
            target_class_score, 3
        ),
        "target_class": target_class,

        "best_hit_description": best_description,
    }

    return row

def build_evidence_rows(
    blast_hits
):

    rows = []

    for gene_id, hits in blast_hits.items():

        limited_hits = sorted(
            hits,
            key=lambda x: (
                x["bitscore"],
                x["pident"],
                x["qcov"],
                -math.log10(
                    max(x["evalue"], 1e-300)
                )
            ),
            reverse=True
        )[:MAX_HITS_PER_GENE]

        sorted_hits = sorted(
            limited_hits,
            key=lambda x: (
                qualify_blast_hit(x),
                x["bitscore"]
            ),
            reverse=True
        )

        for rank, hit in enumerate(
            sorted_hits,
            start=1
        ):

            rows.append({
                "gene_id": gene_id,
                "hit_rank": rank,
                "sseqid": hit["sseqid"],
                "pident": hit["pident"],
                "qcov": hit["qcov"],
                "evalue": hit["evalue"],
                "bitscore": hit["bitscore"],
                "length": hit["length"],
                "qlen": hit["qlen"],
                "slen": hit["slen"],
                "qualifies":
                    "YES"
                    if qualify_blast_hit(hit)
                    else "NO",
                "strong_hit":
                    "YES"
                    if strong_blast_hit(hit)
                    else "NO",
                "stitle": hit["stitle"],
            })

    return rows

def write_rows_to_sheet(ws, rows, columns):

    # Header
    for col_num, column in enumerate(columns, start=1):

        cell = ws.cell(
            row=1,
            column=col_num,
            value=column
        )

        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # Data
    for row_num, row in enumerate(rows, start=2):

        for col_num, column in enumerate(columns, start=1):

            value = row.get(column, "")
            ws.cell(
                row=row_num,
                column=col_num,
                value="" if value is None else value
            )

    ws.freeze_panes = "A2"

    if rows:
        ws.auto_filter.ref = ws.dimensions

    # Column widths
    for col_num, column in enumerate(columns, start=1):

        max_len = len(str(column))

        for row_num in range(
            2,
            min(ws.max_row, 501) + 1
        ):
            value = ws.cell(
                row=row_num,
                column=col_num
            ).value

            if value is not None:
                max_len = max(
                    max_len,
                    len(str(value))
                )

        ws.column_dimensions[
            get_column_letter(col_num)
        ].width = min(
            max(max_len + 2, 12),
            45
        )

    # Wrap text
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


def write_excel(
    path,
    candidates,
    evidence,
    rejected,
    summary_text
):

    if not OPENPYXL_AVAILABLE:
        return False

    workbook = Workbook()
    workbook.remove(workbook.active)

    sheets = [
        ("Candidates", candidates, CANDIDATE_COLUMNS),
        ("BLAST_Evidence", evidence, EVIDENCE_COLUMNS),
        ("Rejected", rejected, REJECTED_COLUMNS),
    ]

    for sheet_name, rows, columns in sheets:

        ws = workbook.create_sheet(sheet_name)

        write_rows_to_sheet(
            ws,
            rows,
            columns
        )

    # Summary
    ws = workbook.create_sheet("Summary")

    for row_num, line in enumerate(
        summary_text.splitlines(),
        start=1
    ):
        ws.cell(
            row=row_num,
            column=1,
            value=line
        )

    ws.column_dimensions["A"].width = 110

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    workbook.save(path)

    return True

def make_summary(
    input_blast,
    input_interpro,
    input_taxonomy,
    input_fasta,
    sequences,
    blast_hits,
    candidates,
    rejected,
    total_blast_rows
):

    keep = sum(row["decision"] == "KEEP" for row in candidates)
    review = sum(row["decision"] == "REVIEW" for row in candidates)

    sequence_stats = validate_sequences(sequences)

    lines = [
        "METISA PLANA TARGET MINING SUMMARY",
        "=" * 60,
        "",
        "INPUT FILES",
        f"BLAST: {input_blast}",
        f"InterPro: {input_interpro}",
        f"Taxonomy: {input_taxonomy}",
        f"FASTA: {input_fasta}",
        "",
        "INPUT STATISTICS",
        f"FASTA proteins: {len(sequences)}",
        f"BLAST rows: {total_blast_rows}",
        f"BLAST query proteins: {len(blast_hits)}",
        f"Sequence empty: {sequence_stats['empty']}",
        f"Internal stop: {sequence_stats['internal_stop']}",
        f"Invalid amino-acid characters: {sequence_stats['invalid_characters']}",
        "",
        "SCORING",
        f"Homology weight: {WEIGHT_HOMOLOGY}%",
        f"Taxonomy weight: {WEIGHT_TAXONOMY}%",
        f"Annotation weight: {WEIGHT_ANNOTATION}%",
        f"InterPro weight: {WEIGHT_INTERPRO}%",
        f"Target class weight: {WEIGHT_TARGET_CLASS}%",
        "",
        "BLAST FILTER",
        f"Identity >= {MIN_IDENTITY_MINING}%",
        f"Query coverage >= {MIN_QUERY_COVERAGE_MINING}%",
        f"E-value <= {MAX_EVALUE_MINING}",
        "",
        "RESULTS",
        f"Final candidates: {len(candidates)}",
        f"KEEP: {keep}",
        f"REVIEW: {review}",
        f"REJECTED: {len(rejected)}",
    ]

    return "\n".join(lines) + "\n"

def main():

    print_header("METISA PLANA TARGET MINING")
    print("Combines BLAST, taxonomy, InterPro and sequence evidence.")

    print_header("INPUT FILES")

    blast_path = ask_existing_file("Filtered blast path: ")
    interpro_path = ask_existing_file("InterProScan TSV path: ")
    taxonomy_path = ask_existing_file("Taxonomy TSV path: ")
    fasta_path = ask_existing_file("Filtered fasta path: ")

    print_header("OUTPUT")

    output_base = ask_output_basename()
    max_candidates = ask_max_candidates()
    output_dir = os.path.dirname(blast_path)

    candidate_path = os.path.join(
        output_dir, output_base + "_candidates.tsv"
    )
    evidence_path = os.path.join(
        output_dir, output_base + "_blast.tsv"
    )
    rejected_path = os.path.join(
        output_dir, output_base + "_rejected.tsv"
    )
    summary_path = os.path.join(
        output_dir, output_base + "_summary.txt"
    )
    excel_path = os.path.join(
        output_dir, output_base + ".xlsx"
    )

    print("\nOutput files:")
    for path in [
        candidate_path,
        evidence_path,
        rejected_path,
        excel_path,
        summary_path
    ]:
        print(" ", path)

    print_header("READING FASTA")

    sequences = read_fasta(fasta_path)
    sequence_stats = validate_sequences(sequences)

    print(f"Proteins loaded: {len(sequences):,}")
    print(f"Empty sequences: {sequence_stats['empty']:,}")
    print(f"Internal stops: {sequence_stats['internal_stop']:,}")
    print(
        f"Invalid characters: "
        f"{sequence_stats['invalid_characters']:,}"
    )

    print_header("READING BLAST")

    blast_hits, total_blast_rows = parse_blast(blast_path)

    print(f"BLAST rows: {total_blast_rows:,}")
    print(f"BLAST queries: {len(blast_hits):,}")

    print_header("READING INTERPRO")

    interpro = parse_interpro(interpro_path)

    print(
        f"Proteins with InterPro annotations: "
        f"{len(interpro):,}"
    )

    print_header("READING TAXONOMY")

    taxonomy = parse_taxonomy(taxonomy_path)

    print(f"Taxonomy records: {len(taxonomy):,}")

    print_header("CALCULATING CANDIDATE SCORES")

    gene_ids = {
        gene_id
        for gene_id, hits in blast_hits.items()
        if gene_id in sequences
        and any(qualify_blast_hit(hit) for hit in hits)
    }

    print(
        f"Proteins passing BLAST filter: "
        f"{len(gene_ids):,}"
    )

    all_candidates = []

    for index, gene_id in enumerate(gene_ids, start=1):

        all_candidates.append(
            build_candidate(
                gene_id,
                blast_hits.get(gene_id, []),
                interpro.get(gene_id, {}),
                taxonomy,
                sequences
            )
        )

        if index % 1000 == 0:
            print(f"Processed {index:,} proteins...")

    candidates = [
        row for row in all_candidates
        if row["decision"] in {"KEEP", "REVIEW"}
    ]

    candidates.sort(
        key=lambda row: (
            row["final_score"],
            row["homology_score"],
            row["taxonomy_score"]
        ),
        reverse=True
    )

    candidates = candidates[:max_candidates]

    for rank, row in enumerate(candidates, start=1):
        row["rank"] = rank

    rejected = [
        row for row in all_candidates
        if row["decision"] == "REJECT"
    ]

    rejected.sort(
        key=lambda row: row["final_score"],
        reverse=True
    )

    evidence_rows = build_evidence_rows(blast_hits)

    summary_text = make_summary(
        blast_path,
        interpro_path,
        taxonomy_path,
        fasta_path,
        sequences,
        blast_hits,
        candidates,
        rejected,
        total_blast_rows
    )

    print_header("WRITING OUTPUT")

    write_tsv(candidate_path, candidates, CANDIDATE_COLUMNS)
    print("Created:", candidate_path)

    write_tsv(evidence_path, evidence_rows, EVIDENCE_COLUMNS)
    print("Created:", evidence_path)

    write_tsv(rejected_path, rejected, REJECTED_COLUMNS)
    print("Created:", rejected_path)

    write_text(summary_path, summary_text)
    print("Created:", summary_path)

    if OPENPYXL_AVAILABLE:

        try:
            write_excel(
                excel_path,
                candidates,
                evidence_rows,
                rejected,
                summary_text
            )
            print("Created:", excel_path)

        except Exception as error:
            print("WARNING: Excel creation failed:")
            print(error)

    else:
        print("WARNING: openpyxl is not installed.")
        print("Excel file was not created.")
        print("Install with: pip install openpyxl")

    print_header("DONE")

    keep_count = sum(
        row["decision"] == "KEEP"
        for row in candidates
    )
    review_count = sum(
        row["decision"] == "REVIEW"
        for row in candidates
    )

    print(f"Top candidates exported: {len(candidates):,}")
    print(f"KEEP: {keep_count:,}")
    print(f"REVIEW: {review_count:,}")
    print(f"Rejected: {len(rejected):,}")
    print(f"\nAll outputs use the same base name: {output_base}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nProcess cancelled by user.")
    except Exception as error:
        print("\n" + "=" * 78)
        print("ERROR")
        print("=" * 78)
        print(error)